"""Import of borehole logs from CSV or Excel (.xlsx), and the template for them.

One flat table, one record per row (Excel: the sheet "Soil data", else the first sheet; the
header is the first row with a "Borehole" column):

- ``Borehole`` (required): the borehole's id, e.g. BH-1. ``X``, ``Y``, ``Ground level``,
  ``Water table`` (depth below the ground) and ``Borehole depth`` belong to the borehole: give
  them on any of its rows (the first value counts; a different one later is an error).
- ``Record``: ``layer``, ``spt`` or ``sample``. When empty, a row with Top and Bottom is a
  layer, a row with N (or N1, N2) an SPT record, and any other row with a Depth a sample.
- Layers: ``Top``, ``Bottom`` (depths below the ground), ``Description``, ``Soil type`` (clay,
  silt, sand, gravel, organic, fill, rock; when empty, from the USCS group or the last soil
  word of the description), ``USCS`` and the properties below.
- SPT records: ``Depth``, ``N`` (field blow count) or the increments ``Seat``, ``N1``, ``N2``
  (blows per 150 mm; N = N1 + N2), ``Sample no``, ``Energy ratio`` (%, optional). A sheet
  that numbers the three increments ``N1``, ``N2``, ``N3`` is read with N1 as the seating drive
  and N = N2 + N3 (a note says so).
- Samples: ``Depth``, ``Sample type``, ``Sample no`` and the properties below.
- Properties: ``Unit weight`` and ``Saturated unit weight`` (kN/m³), ``cu`` and ``c`` (kPa),
  ``phi`` (°), ``E`` (MPa), ``Fines``, ``PI``, ``LL``, ``PL`` and ``Moisture`` (%), ``Cc``,
  ``Cr``, ``e0``, ``OCR``. Each record type takes those of the file format it has (a layer
  has no plastic limit, a sample no Cr); the others are reported.

Lengths are in metres unless the header names the unit: ``Top (mm)``, ``Depth (ft)``; m, cm,
mm and ft are understood. The other columns are in the units above unless the header names
another, which is converted: unit weight from g/cc, t/m³, kg/m³ or pcf (a density times g);
cu, c from kg/cm², t/m², tsf, psi, psf, MPa; E from kPa, kg/cm², tsf, psi. A header unit that
is not understood is reported, and a "Bulk density" column must name its unit. A unit weight
below 5 kN/m³ is reported as a probable density. Header case and spaces do not matter.

Problems are reported by row number (as the spreadsheet shows it) and column.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from saengine.library.soil import SOIL_TYPES, SoilReport

_LENGTH_UNITS = {"m": 1000.0, "mm": 1.0, "cm": 10.0, "ft": 304.8}

# Factors from a header's unit to the column's unit (g = 9.80665 m/s²). A density (mass per
# volume) is turned into a unit weight. Pressures: 1 kg/cm² = 98.0665 kPa, 1 tsf (short ton per
# ft²) = 95.76 kPa, 1 t/m² = 9.80665 kPa, 1 psi = 6.894757 kPa, 1 psf = 0.04788 kPa.
_G = 9.80665
_PRESSURE_KPA = {
    "kpa": 1.0,
    "kn/m2": 1.0,
    "mpa": 1000.0,
    "kg/cm2": _G * 10,
    "kgf/cm2": _G * 10,
    "ksc": _G * 10,
    "t/m2": _G,
    "tsf": 95.76,
    "psi": 6.894757,
    "psf": 0.04788026,
    "ksf": 47.88026,
}
_UNITS: dict[str, dict[str, float]] = {
    "kN/m³": {
        "kn/m3": 1.0,
        "g/cc": _G,
        "g/cm3": _G,
        "t/m3": _G,
        "mg/m3": _G,
        "kg/m3": _G / 1000,
        "pcf": 0.1570875,
        "lb/ft3": 0.1570875,
    },
    "kPa": _PRESSURE_KPA,
    "MPa": {k: v / 1000 for k, v in _PRESSURE_KPA.items()},
    "°": {"°": 1.0, "deg": 1.0, "degree": 1.0, "degrees": 1.0},
    "%": {"%": 1.0, "percent": 1.0},
    # Counts and ratios: only descriptive units are accepted.
    "": {"-": 1.0, "blows": 1.0, "blows/300mm": 1.0, "blows/ft": 1.0, "bpf": 1.0},
}
_DENSITY_NAMES = {"bulkdensity", "density"}


def _unit_key(unit: str) -> str:
    return (
        unit.lower()
        .replace(" ", "")
        .replace("³", "3")
        .replace("²", "2")
        .replace("^", "")
        .replace("cu.m", "m3")
    )


def _unit_scale(column: Column, unit: str | None, name: str) -> tuple[float | None, str | None]:
    """(factor to the column's unit, problem) for a header's unit."""
    if column.unit == "length":
        scale = _LENGTH_UNITS.get(_unit_key(unit or "m"))
        return (
            (scale, None)
            if scale
            else (None, f"unknown length unit '{unit}' (use m, cm, mm or ft)")
        )
    if column.key in _TEXT:
        return 1.0, None
    if not unit:
        if name in _DENSITY_NAMES:
            return None, "give the unit of the density in the header, e.g. 'Bulk density (g/cc)'"
        return 1.0, None
    table = _UNITS[column.unit or ""]
    scale = table.get(_unit_key(unit))
    if scale is None:
        known = ", ".join(table)
        return None, f"unknown unit '{unit}' (use {column.unit or 'no unit'}" + (
            f"; or {known})" if known else ")"
        )
    return scale, None


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    unit: str | None  # "length" for lengths, else the display unit
    help: str
    aliases: tuple[str, ...] = ()


COLUMNS: tuple[Column, ...] = (
    Column(
        "borehole",
        "Borehole",
        None,
        "Borehole id, e.g. BH-1 (every row)",
        ("bh", "boreholeid", "boreholeno", "id"),
    ),
    Column(
        "record", "Record", None, "layer, spt or sample (empty: inferred)", ("recordtype", "kind")
    ),
    Column("x", "X", "length", "Plan position of the borehole, project coordinates", ("easting",)),
    Column("y", "Y", "length", "Plan position of the borehole, project coordinates", ("northing",)),
    Column(
        "lat",
        "Latitude",
        "°",
        "Borehole: WGS 84 latitude; with the longitude it places the borehole",
        ("latitude",),
    ),
    Column("lon", "Longitude", "°", "Borehole: WGS 84 longitude", ("longitude", "long", "lng")),
    Column(
        "btmEasting",
        "BTM easting",
        None,
        "Borehole: Bangladesh Transverse Mercator easting [m], as survey drawings give it",
        ("btme",),
    ),
    Column(
        "btmNorthing",
        "BTM northing",
        None,
        "Borehole: Bangladesh Transverse Mercator northing [m]",
        ("btmn",),
    ),
    Column("date", "Date", None, "Borehole: date drilled", ("datedrilled", "dateofboring")),
    Column(
        "method",
        "Method",
        None,
        "Borehole: boring method, e.g. wash boring",
        ("boringmethod", "drillingmethod"),
    ),
    Column(
        "groundLevel",
        "Ground level",
        "length",
        "Elevation of the ground at the borehole (project datum)",
        ("gl", "rl", "elevation", "existinggroundlevel"),
    ),
    Column(
        "waterTable",
        "Water table",
        "length",
        "Depth of the ground water below the ground",
        ("wt", "gwt", "gwl", "watertabledepth", "groundwater"),
    ),
    Column(
        "boreholeDepth",
        "Borehole depth",
        "length",
        "Termination depth of the borehole",
        ("termination", "terminationdepth", "totaldepth"),
    ),
    Column(
        "top",
        "Top",
        "length",
        "Layer: depth of the top below the ground",
        ("from", "depthfrom", "topdepth"),
    ),
    Column(
        "bottom",
        "Bottom",
        "length",
        "Layer: depth of the bottom below the ground",
        ("to", "depthto", "bottomdepth"),
    ),
    Column(
        "depth",
        "Depth",
        "length",
        "SPT record or sample: depth below the ground",
        ("sptdepth", "sampledepth"),
    ),
    Column(
        "description", "Description", None, "Layer: description from the log", ("soildescription",)
    ),
    Column(
        "soilType",
        "Soil type",
        None,
        "Layer: clay, silt, sand, gravel, organic, fill or rock",
        ("soil", "type"),
    ),
    Column("uscs", "USCS", None, "Layer: USCS group, e.g. CL, ML, SM", ("uscsgroup", "group")),
    Column(
        "N",
        "N",
        None,
        "SPT: field blow count for 300 mm, uncorrected",
        ("spt", "sptn", "nvalue", "blows", "fieldn"),
    ),
    Column(
        "depthTo",
        "SPT to",
        "length",
        "SPT record: depth of the end of the drive (empty: depth + 0.45 m)",
        ("driveto",),
    ),
    Column("seat", "Seat", None, "SPT: blows for the first (seating) 150 mm", ("seating", "n0")),
    Column("N1", "N1", None, "SPT: blows for the second 150 mm", ("blows2",)),
    Column("N2", "N2", None, "SPT: blows for the third 150 mm", ("blows3",)),
    Column(
        "sampleNumber",
        "Sample no",
        None,
        "SPT record or sample: the sample number, e.g. D-3",
        ("sampleno", "samplenumber", "sampleid"),
    ),
    Column("energyRatio", "Energy ratio", "%", "SPT: hammer energy ratio of this test", ("er",)),
    Column(
        "refusal",
        "Refusal",
        None,
        "SPT: R when the drive stopped at 50 blows, or blows/penetration as 50/100 (mm)",
        ("ref",),
    ),
    Column(
        "recovery",
        "Recovery",
        "%",
        "SPT record or sample: length recovered over the length driven",
        ("rec",),
    ),
    Column("sampleType", "Sample type", None, "Sample: e.g. SPT, UD", ("sample",)),
    Column(
        "unitWeight",
        "Unit weight",
        "kN/m³",
        "Layer or sample: unit weight (bulk)",
        ("gamma", "γ", "bulkunitweight", "bulkdensity", "density"),
    ),
    Column(
        "saturatedUnitWeight",
        "Saturated unit weight",
        "kN/m³",
        "Layer: saturated unit weight",
        ("gammasat", "γsat"),
    ),
    Column("cu", "cu", "kPa", "Undrained shear strength", ("su", "undrainedshearstrength")),
    Column("c", "c", "kPa", "Effective cohesion", ("cohesion",)),
    Column(
        "phi", "phi", "°", "Effective friction angle", ("φ", "frictionangle", "angleoffriction")
    ),
    Column("E", "E", "MPa", "Layer: Young's modulus", ("modulus", "youngsmodulus")),
    Column(
        "finesContent",
        "Fines",
        "%",
        "Fines content, passing 75 µm",
        ("fc", "finescontent", "passing200"),
    ),
    Column("plasticityIndex", "PI", "%", "Layer: plasticity index", ("plasticityindex", "ip")),
    Column("liquidLimit", "LL", "%", "Liquid limit", ("liquidlimit",)),
    Column("plasticLimit", "PL", "%", "Sample: plastic limit", ("plasticlimit", "wp")),
    Column(
        "moisture",
        "Moisture",
        "%",
        "Sample: natural moisture content",
        ("w", "moisturecontent", "watercontent", "mc"),
    ),
    Column("Cc", "Cc", None, "Compression index", ("compressionindex",)),
    Column(
        "Cr",
        "Cr",
        None,
        "Layer: recompression index",
        ("cs", "recompressionindex", "swellingindex"),
    ),
    Column("e0", "e0", None, "Initial void ratio", ("voidratio", "e")),
    Column("OCR", "OCR", None, "Layer: overconsolidation ratio", ("overconsolidationratio",)),
)

_BOREHOLE_FIELDS = (
    "x",
    "y",
    "lat",
    "lon",
    "btmEasting",
    "btmNorthing",
    "groundLevel",
    "waterTable",
    "boreholeDepth",
)
#: The borehole's text values (the first row that gives one; another value is a problem).
_BOREHOLE_TEXT = ("date", "method")
#: The file format's names of the borehole values the template names otherwise.
_BOREHOLE_KEYS = {"boreholeDepth": "depth", "btmEasting": "easting", "btmNorthing": "northing"}
_LAYER_FIELDS = (
    "unitWeight",
    "saturatedUnitWeight",
    "cu",
    "c",
    "phi",
    "E",
    "finesContent",
    "plasticityIndex",
    "liquidLimit",
    "Cc",
    "Cr",
    "e0",
    "OCR",
)
_SAMPLE_FIELDS = (
    "moisture",
    "unitWeight",
    "cu",
    "c",
    "phi",
    "finesContent",
    "liquidLimit",
    "plasticLimit",
    "Cc",
    "e0",
)
_TEXT = {
    "borehole",
    "record",
    "description",
    "soilType",
    "uscs",
    "sampleType",
    "sampleNumber",
    "date",
    "method",
    "refusal",
}

_WORDS = {
    "clay": "clay",
    "silt": "silt",
    "sand": "sand",
    "gravel": "gravel",
    "peat": "organic",
    "organic": "organic",
    "fill": "fill",
    "rock": "rock",
}


def _norm(text: str) -> str:
    return re.sub(r"[^0-9a-zμγφ]", "", text.lower())


def _parse_header(text: str) -> tuple[str, str | None]:
    """(normalized name, unit) of a header cell such as 'Top (mm)'."""
    unit = None
    match = re.search(r"[\(\[]\s*([^\)\]]*)\s*[\)\]]", text)
    if match:
        unit = match.group(1).strip().lower()
        text = text[: match.start()] + text[match.end() :]
    return _norm(text), unit


_BY_NAME: dict[str, Column] = {}
for _column in COLUMNS:
    for _name in (_column.key, _column.header, *_column.aliases):
        _BY_NAME.setdefault(_norm(_name), _column)


@dataclass
class _Borehole:
    id: str
    values: dict[str, tuple[float, int]] = field(default_factory=dict)
    text: dict[str, tuple[str, int]] = field(default_factory=dict)
    layers: list[tuple[int, dict[str, Any]]] = field(default_factory=list)
    spt: list[tuple[int, dict[str, Any]]] = field(default_factory=list)
    samples: list[tuple[int, dict[str, Any]]] = field(default_factory=list)


def _from_uscs(group: str) -> str | None:
    """Soil type of a USCS group symbol (the first of a dual symbol such as SM-SC)."""
    code = group.strip().upper()[:2]
    if code == "PT" or code[:1] == "O":
        return "organic"
    return (
        {"C": "clay", "M": "silt", "S": "sand", "G": "gravel"}.get(code[:1])
        if len(code) == 2
        else None
    )


def soil_type_of(value: str | None, uscs: str | None, description: str | None) -> str | None:
    """The file format's soil type from the cell, else the USCS group, else the description's
    last soil noun ("Grey silty fine SAND" is sand; adjectives such as silty are skipped)."""
    if value:
        key = value.strip().lower()
        if key in SOIL_TYPES:
            return key
        if key in _WORDS:
            return _WORDS[key]
        if key in ("made ground", "madeground"):
            return "fill"
        return _from_uscs(key) if re.fullmatch(r"[a-z]{2}(-[a-z]{2})?", key) else None
    if uscs and (kind := _from_uscs(uscs)):
        return kind
    if description:
        for word in reversed(re.findall(r"[a-z]+", description.lower())):
            for stem, kind in _WORDS.items():
                if word in (stem, stem + "s"):
                    return kind
    return None


def _rows_csv(path: Path) -> list[tuple[int, list[str]]]:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [(i + 1, row) for i, row in enumerate(csv.reader(io.StringIO(text), dialect))]


def _rows_xlsx(path: Path) -> list[tuple[int, list[Any]]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        names = {name.lower(): name for name in book.sheetnames}
        sheet = book[names.get("soil data", book.sheetnames[0])]
        return [
            (i + 1, ["" if v is None else v for v in row])
            for i, row in enumerate(sheet.iter_rows(values_only=True))
        ]
    finally:
        book.close()


def read_rows(path: str | Path) -> list[tuple[int, list[Any]]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv" or suffix == ".txt":
        return _rows_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _rows_xlsx(path)
    raise ValueError(f"{path.name}: use a .csv or .xlsx file")


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"'{value}' is not a number")  # noqa: TRY004 - reported as a cell problem
    if isinstance(value, int | float):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text in ("-", "—", "n/a", "na"):
        return None
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"'{value}' is not a number") from None


def parse_rows(
    rows: list[tuple[int, list[Any]]],
) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    """Boreholes (file-format dicts) from table rows, the problems and notes."""
    problems: list[str] = []
    notes: list[str] = []
    header_at = None
    for index, (number, cells) in enumerate(rows):
        names = [_parse_header(str(c))[0] for c in cells]
        if any(_BY_NAME.get(n) is not None and _BY_NAME[n].key == "borehole" for n in names):
            header_at = index
            break
    if header_at is None:
        return [], ["No header row with a Borehole column (see the template)."], []
    header_number, header = rows[header_at]
    columns: list[tuple[Column, float] | None] = []
    ignored = []
    # Log sheets that number all three 150 mm increments N1, N2, N3 count the seating drive as
    # N1: their N is N2 + N3. Map them onto seat, N1, N2 (this file format's names).
    names = {_parse_header(str(c))[0] for c in header}
    renumber = {}
    if "n3" in names:
        seat_given = any(_BY_NAME.get(n) is not None and _BY_NAME[n].key == "seat" for n in names)
        if seat_given:
            problems.append(
                f"row {header_number}: both a seating column and N1, N2, N3: with N3 the "
                "increments are read as N1 = seating, N2 and N3 = N; remove one of them"
            )
            return [], problems, notes
        renumber = {"n1": "seat", "n2": "n1", "n3": "n2"}
        notes.append(
            "The increments are numbered N1, N2, N3: N1 is read as the seating 150 mm and "
            "N = N2 + N3."
        )
    for cell in header:
        name, unit = _parse_header(str(cell))
        name = renumber.get(name, name)
        column = _BY_NAME.get(name)
        if column is None:
            if name:
                ignored.append(str(cell))
            columns.append(None)
            continue
        scale, problem = _unit_scale(column, unit, name)
        if problem or scale is None:
            problems.append(f"row {header_number}: {cell}: {problem}")
            columns.append(None)
            continue
        columns.append((column, scale))
    if ignored:
        notes.append(f"Columns not recognized, ignored: {', '.join(ignored)}.")
    boreholes: dict[str, _Borehole] = {}
    for number, cells in rows[header_at + 1 :]:
        values: dict[str, Any] = {}
        bad: list[str] = []
        for spec, cell in zip(columns, cells, strict=False):
            if spec is None or cell is None or (isinstance(cell, str) and not cell.strip()):
                continue
            column, scale = spec
            if column.key in _TEXT:
                # Excel gives a date cell as a datetime: the date is what the log says.
                text = cell.strftime("%Y-%m-%d") if hasattr(cell, "strftime") else str(cell)
                values[column.key] = text.strip()
                continue
            try:
                number_value = _number(cell)
            except ValueError as exc:
                bad.append(f"{column.header}: {exc}")
                continue
            if number_value is not None:
                values[column.key] = number_value if scale == 1 else round(number_value * scale, 4)
                if column.key == "unitWeight" and 0 < values[column.key] < 5:
                    bad.append(
                        f"{column.header}: {values[column.key]:g} kN/m³ is too low for a unit "
                        "weight: a density in g/cc or t/m³? Name the unit in the header, e.g. "
                        "'Bulk density (g/cc)'"
                    )
        if not values and not bad:
            continue
        bh_id = values.get("borehole")
        if not bh_id:
            problems.append(f"row {number}: no borehole id")
            continue
        where = f"row {number} ({bh_id})"
        if bad:
            problems += [f"{where}: {message}" for message in bad]
            continue
        hole = boreholes.setdefault(bh_id, _Borehole(bh_id))
        for key in _BOREHOLE_FIELDS:
            if key in values:
                if key in hole.values and abs(hole.values[key][0] - values[key]) > 1e-6:
                    first = hole.values[key][1]
                    problems.append(
                        f"{where}: {_column_named(key).header}: {values[key]:g} differs from "
                        f"{hole.values[key][0]:g} on row {first}"
                    )
                else:
                    hole.values.setdefault(key, (values[key], number))
        for key in _BOREHOLE_TEXT:
            if key in values:
                if key in hole.text and hole.text[key][0] != values[key]:
                    problems.append(
                        f"{where}: {_column_named(key).header}: '{values[key]}' differs from "
                        f"'{hole.text[key][0]}' on row {hole.text[key][1]}"
                    )
                else:
                    hole.text.setdefault(key, (values[key], number))
        record = (values.get("record") or "").strip().lower()
        if not record:
            if "top" in values or "bottom" in values:
                record = "layer"
            elif "N" in values or "N1" in values or "N2" in values or "refusal" in values:
                record = "spt"
            elif "depth" in values:
                record = "sample"
            else:
                continue  # a row with the borehole's own values only
        if record in ("layer", "stratum", "l"):
            _layer(hole, number, where, values, problems, notes)
        elif record in ("spt", "s", "n"):
            _spt(hole, number, where, values, problems)
        elif record in ("sample", "lab", "test"):
            _sample(hole, number, where, values, problems)
        else:
            problems.append(f"{where}: Record: '{values['record']}' is not layer, spt or sample")
    result = []
    for hole in boreholes.values():
        data: dict[str, Any] = {"id": hole.id}
        for key in _BOREHOLE_FIELDS:
            if key in hole.values:
                data[_BOREHOLE_KEYS.get(key, key)] = hole.values[key][0]
        for key in _BOREHOLE_TEXT:
            if key in hole.text:
                data[key] = hole.text[key][0]
        data["layers"] = [v for _, v in sorted(hole.layers, key=lambda r: r[1]["top"])]
        data["spt"] = [v for _, v in sorted(hole.spt, key=lambda r: r[1]["depth"])]
        data["samples"] = [v for _, v in sorted(hole.samples, key=lambda r: r[1]["depth"])]
        result.append(data)
    return result, problems, notes


def _column_named(key: str) -> Column:
    return next(c for c in COLUMNS if c.key == key)


def _extra(values: dict[str, Any], allowed: tuple[str, ...], kind: str) -> list[str]:
    own = {"borehole", "record", *_BOREHOLE_FIELDS, *_BOREHOLE_TEXT}
    by_kind = {
        "layer": {"top", "bottom", "description", "soilType", "uscs"},
        "spt": {
            "depth",
            "depthTo",
            "N",
            "seat",
            "N1",
            "N2",
            "sampleNumber",
            "energyRatio",
            "refusal",
            "recovery",
        },
        "sample": {"depth", "sampleType", "sampleNumber", "recovery"},
    }[kind]
    return [
        _column_named(k).header
        for k in values
        if k not in own and k not in by_kind and k not in allowed
    ]


def _layer(
    hole: _Borehole,
    number: int,
    where: str,
    values: dict[str, Any],
    problems: list[str],
    notes: list[str],
) -> None:
    missing = [name for key, name in (("top", "Top"), ("bottom", "Bottom")) if key not in values]
    if missing:
        problems.append(f"{where}: a layer needs {' and '.join(missing)}")
        return
    kind = soil_type_of(values.get("soilType"), values.get("uscs"), values.get("description"))
    if kind is None:
        problems.append(
            f"{where}: Soil type: '{values.get('soilType', '')}' is not one of "
            f"{', '.join(SOIL_TYPES)} (and neither the USCS group nor the description tell)"
        )
        return
    if not values.get("soilType"):
        notes.append(f"{where}: soil type '{kind}' taken from the USCS group or the description")
    layer: dict[str, Any] = {
        "top": values["top"],
        "bottom": values["bottom"],
        "description": values.get("description", ""),
        "soilType": kind,
    }
    if values.get("uscs"):
        layer["uscs"] = values["uscs"]
    for key in _LAYER_FIELDS:
        if key in values:
            layer[key] = values[key]
    extra = _extra(values, _LAYER_FIELDS, "layer")
    if extra:
        problems.append(f"{where}: a layer has no {', '.join(extra)} (give it on a sample row)")
    hole.layers.append((number, layer))


def _spt(
    hole: _Borehole, number: int, where: str, values: dict[str, Any], problems: list[str]
) -> None:
    has_blows = "N" in values or "N1" in values or "N2" in values or "refusal" in values
    if "depth" not in values or not has_blows:
        problems.append(f"{where}: an SPT record needs Depth and N (or N1 and N2, or a refusal)")
        return
    record: dict[str, Any] = {"depth": values["depth"]}
    if "depthTo" in values:
        record["depthTo"] = values["depthTo"]
    if "refusal" in values:
        refusal = _refusal(values["refusal"])
        if refusal is None:
            problems.append(
                f"{where}: Refusal: '{values['refusal']}' is not R (or yes) nor blows/penetration "
                "as 50/100"
            )
            return
        if refusal is not False:
            record["refusal"] = refusal
    for key in ("N", "seat", "N1", "N2"):
        if key in values:
            n = values[key]
            if n != int(n):
                header = _column_named(key).header
                problems.append(f"{where}: {header}: {n:g} is not a whole number of blows")
                return
            record[key] = int(n)
    if ("N1" in record) != ("N2" in record) and "refusal" not in record:
        problems.append(f"{where}: give both N1 and N2")
        return
    if "N1" in record and "N" in record and record["N"] != record["N1"] + record["N2"]:
        problems.append(
            f"{where}: N {record['N']} is not N1 + N2 = {record['N1']} + {record['N2']}"
        )
        return
    if values.get("sampleNumber"):
        record["sampleNumber"] = values["sampleNumber"]
    if "energyRatio" in values:
        record["energyRatio"] = values["energyRatio"]
    if "recovery" in values:
        record["recovery"] = values["recovery"]
    extra = _extra(values, (), "spt")
    if extra:
        problems.append(
            f"{where}: an SPT record has no {', '.join(extra)} (give lab values on a sample row)"
        )
    hole.spt.append((number, record))


def _refusal(text: str) -> dict[str, Any] | bool | None:
    """A refusal cell: True for R (yes, x), blows and penetration for "50/100" (mm), False for
    no (empty, no, -); None when it cannot be read."""
    word = text.strip().lower()
    if word in ("", "no", "n", "-", "0", "false"):
        return False
    if word in ("r", "ref", "refusal", "y", "yes", "x", "1", "true"):
        return True
    match = re.fullmatch(r"(\d+)\s*/\s*(\d+(?:\.\d+)?)\s*(mm)?", word)
    if match:
        return {"blows": int(match.group(1)), "penetration": float(match.group(2))}
    return None


def _sample(
    hole: _Borehole, number: int, where: str, values: dict[str, Any], problems: list[str]
) -> None:
    if "depth" not in values:
        problems.append(f"{where}: a sample needs a Depth")
        return
    sample: dict[str, Any] = {"depth": values["depth"]}
    if values.get("sampleType"):
        sample["type"] = values["sampleType"]
    if values.get("sampleNumber"):
        sample["number"] = values["sampleNumber"]
    if "recovery" in values:
        sample["recovery"] = values["recovery"]
    for key in _SAMPLE_FIELDS:
        if key in values:
            sample[key] = values[key]
    extra = _extra(values, _SAMPLE_FIELDS, "sample")
    if extra:
        problems.append(f"{where}: a sample has no {', '.join(extra)} (give it on a layer row)")
    hole.samples.append((number, sample))


def import_soil(path: str, soil_text: str | None = None) -> dict[str, Any]:
    """The soil report with the boreholes of a CSV or Excel file: ``{"text", "notes"}`` (a
    borehole with an id already in the report replaces it), or ``{"problems"}``. The rows of
    the template's example borehole (:data:`EXAMPLE_ID`) are skipped.

    The editor asks with an empty report and merges the boreholes itself, asking the engineer
    what to do with an id the report already has (replace, keep both, skip)."""
    from saengine.soil.report import dumps, messages, parse

    try:
        rows = read_rows(path)
    except Exception as exc:  # noqa: BLE001 - unreadable files are reported to the user
        return {"problems": [f"{Path(path).name}: {exc}"]}
    boreholes, problems, notes = parse_rows(rows)
    if problems:
        return {"problems": problems}
    example = [hole for hole in boreholes if hole["id"] == EXAMPLE_ID]
    boreholes = [hole for hole in boreholes if hole["id"] != EXAMPLE_ID]
    if not boreholes:
        if example:
            return {
                "problems": [
                    (
                        f"{Path(path).name}: only the template's example borehole {EXAMPLE_ID}, "
                        "which is not imported: add your boreholes' rows under it"
                    )
                ]
            }
        return {"problems": [f"{Path(path).name}: no borehole rows below the header"]}
    if example:
        notes.append(f"{EXAMPLE_ID}: the template's example, not imported.")
    base = parse(soil_text) if soil_text else SoilReport()
    if isinstance(base, list):
        return {"problems": base}
    data = base.model_dump(mode="json", by_alias=True, exclude_none=True)
    existing = {b["id"]: i for i, b in enumerate(data["boreholes"])}
    for hole in boreholes:
        counts = (
            f"{len(hole['layers'])} layers, {len(hole['spt'])} SPT records, "
            f"{len(hole['samples'])} samples"
        )
        if hole["id"] in existing:
            data["boreholes"][existing[hole["id"]]] = hole
            notes.append(f"{hole['id']}: replaced ({counts}).")
        else:
            data["boreholes"].append(hole)
            notes.append(f"{hole['id']}: added ({counts}).")
    try:
        report = SoilReport.model_validate(data)
    except ValidationError as exc:
        return {"problems": messages(exc, data, Path(path).name)}
    return {"text": dumps(report), "notes": notes}


# ─────────────────────────────── template ───────────────────────────────

#: The id of the template's example borehole: the importer skips it, so a template filled in
#: beside the example never brings the example in (nor replaces a real BH-1).
EXAMPLE_ID = "EXAMPLE-1"

#: Example rows of the template: a borehole with layers, SPT records and samples (lengths in m).
EXAMPLE: list[dict[str, Any]] = [
    {
        "borehole": EXAMPLE_ID,
        "record": "layer",
        "x": 0,
        "y": 0,
        "date": "2026-01-15",
        "method": "Wash boring",
        "groundLevel": 0,
        "waterTable": 1.5,
        "boreholeDepth": 15,
        "top": 0,
        "bottom": 1.5,
        "description": "Brownish grey clayey silt (fill)",
        "soilType": "fill",
        "unitWeight": 17,
    },
    {
        "borehole": EXAMPLE_ID,
        "record": "layer",
        "top": 1.5,
        "bottom": 6,
        "description": "Soft grey silty clay",
        "soilType": "clay",
        "uscs": "CL",
        "unitWeight": 18,
        "saturatedUnitWeight": 18.5,
        "cu": 25,
        "liquidLimit": 42,
        "plasticityIndex": 18,
        "Cc": 0.28,
        "e0": 0.95,
    },
    {
        "borehole": EXAMPLE_ID,
        "record": "layer",
        "top": 6,
        "bottom": 15,
        "description": "Medium dense grey fine sand",
        "soilType": "sand",
        "uscs": "SM",
        "saturatedUnitWeight": 19.5,
        "finesContent": 18,
    },
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 1.5, "N": 4},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 3.0, "N": 3},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 4.5, "N": 5},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 6.0, "N": 9},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 7.5, "N": 12},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 9.0, "N": 15},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 10.5, "N": 18},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 12.0, "N": 22},
    {"borehole": EXAMPLE_ID, "record": "spt", "depth": 13.5, "N": 25},
    {
        "borehole": EXAMPLE_ID,
        "record": "spt",
        "depth": 15.0,
        "depthTo": 15.45,
        "seat": 12,
        "N1": 13,
        "N2": 15,
        "sampleNumber": "D-10",
        "recovery": 80,
    },
    {
        "borehole": EXAMPLE_ID,
        "record": "sample",
        "depth": 3.0,
        "sampleType": "UD",
        "moisture": 34,
        "unitWeight": 18.2,
        "cu": 24,
        "liquidLimit": 41,
        "plasticLimit": 23,
    },
    {
        "borehole": EXAMPLE_ID,
        "record": "sample",
        "depth": 9.0,
        "sampleType": "SPT",
        "moisture": 24,
        "finesContent": 16,
    },
]


def _header(column: Column) -> str:
    if column.unit == "length":
        return f"{column.header} (m)"
    return f"{column.header} ({column.unit})" if column.unit else column.header


def template_rows() -> list[list[Any]]:
    rows: list[list[Any]] = [[_header(c) for c in COLUMNS]]
    for example in EXAMPLE:
        rows.append([example.get(c.key, "") for c in COLUMNS])
    return rows


def soil_template(path: str) -> dict[str, Any]:
    """Write the import template (.csv or .xlsx by the extension) and return ``{"path"}``."""
    target = Path(path)
    suffix = target.suffix.lower()
    rows = template_rows()
    if suffix == ".csv":
        with target.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        return {"path": str(target)}
    if suffix != ".xlsx":
        return {"problems": [f"{target.name}: the template is a .csv or .xlsx file"]}
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Soil data"
    for row in rows:
        sheet.append(row)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8EEF4")
    for cell in sheet[1]:
        cell.font = bold
        cell.fill = fill
    for index, column in enumerate(COLUMNS, start=1):
        width = max(len(_header(column)), 8 if column.key != "description" else 34)
        sheet.column_dimensions[get_column_letter(index)].width = width + 2
    sheet.freeze_panes = "B2"
    notes = book.create_sheet("Read me")
    notes.append(["SAEngine soil data: one record (layer, SPT record or sample) per row."])
    notes.append(
        ["Lengths in metres unless the header says (mm); depths below the borehole's ground."]
    )
    notes.append(
        [
            (
                "Borehole values (X and Y, or latitude and longitude, or BTM; ground level, "
                "water table, depth, date, method) on any of its rows."
            )
        ]
    )
    notes.append([])
    notes.append(["Column", "Unit", "Meaning"])
    for cell in notes[5]:
        cell.font = bold
    for column in COLUMNS:
        unit = "m" if column.unit == "length" else column.unit or ""
        notes.append([column.header, unit, column.help])
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 8
    notes.column_dimensions["C"].width = 70
    book.save(target)
    return {"path": str(target)}
