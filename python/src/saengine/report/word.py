"""The design report as a Word file (``11 Exports/report.docx``), with python-docx.

Draws the same document as the PDF (:func:`saengine.report.document.report_document`): the cover,
the contents (a Word field, filled in when Word opens the file), each section's blocks and the
landscape appendix. Headings use Word's Heading 1-3 styles, numbered here as the PDF numbers them
(1, 1.1, 1.1.1; the appendix A, A.1), so the engineer can restyle or renumber them in Word.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x5F, 0x5F, 0x5F)
FAIL_INK = RGBColor(0xA1, 0x15, 0x0C)
PASS_INK = RGBColor(0x1F, 0x4D, 0x21)
HEAD_FILL = "E8E8E8"
FAIL_FILL = "F9D9D6"
WARN_FILL = "FDF1D6"
PASS_FILL = "DCEFDC"
RULE = "AFAFAF"

PAPER = {"a4": (Mm(210), Mm(297)), "us-letter": (Mm(215.9), Mm(279.4))}
ALIGN = {
    "l": WD_ALIGN_PARAGRAPH.LEFT,
    "r": WD_ALIGN_PARAGRAPH.RIGHT,
    "c": WD_ALIGN_PARAGRAPH.CENTER,
}
CHECK_HEAD = ["Check", "Demand", "Capacity", "Unit", "Util.", "Combination", "Result"]
CHECK_ALIGN = "lrrlrlc"


def render_docx(
    document: dict[str, Any], project_root: str | Path | None = None
) -> tuple[bytes, list[str]]:
    """The Word file of the report ``document`` and what could not be drawn."""
    writer = _Writer(document, project_root)
    writer.write()
    out = io.BytesIO()
    writer.doc.save(out)
    return out.getvalue(), writer.warnings


class _Writer:
    def __init__(self, document: dict[str, Any], project_root: str | Path | None) -> None:
        self.data = document
        self.root = Path(project_root) if project_root else None
        self.doc = Document()
        self.warnings: list[str] = []
        self.numbers = [0, 0, 0]
        self.appendix = False
        self._styles()

    # ─────────────────────────── set-up ───────────────────────────

    def _styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Cambria"
        normal.font.size = Pt(10)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(4)
        for level, size in ((1, 15), (2, 11.5), (3, 10)):
            style = styles[f"Heading {level}"]
            style.font.name = "Cambria"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = INK
            style.paragraph_format.space_before = Pt(14 if level > 1 else 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        width, height = PAPER.get(self.data.get("paper", "a4"), PAPER["a4"])
        section = self.doc.sections[0]
        section.page_width, section.page_height = width, height
        section.left_margin, section.right_margin = Mm(20), Mm(18)
        section.top_margin, section.bottom_margin = Mm(24), Mm(22)
        # Word fills in the contents (and the page fields) when it opens the file.
        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def _header_footer(self, section: Any) -> None:
        meta = self.data["meta"]
        letterhead = self.data["letterhead"]
        header = section.header
        header.is_linked_to_previous = False
        first = header.paragraphs[0]
        if letterhead.get("enabled"):
            self._letterhead(first, size=7.5)
            first = header.add_paragraph()
        run = first.add_run(f"{meta['project']}\t\t{meta['title']}")
        _font(run, 8, MUTED)
        _bottom_rule(first)
        footer = section.footer
        footer.is_linked_to_previous = False
        line = footer.paragraphs[0]
        _font(line.add_run(f"SAEngine {meta['version']}, {meta['isoDate']}\t\t"), 8, MUTED)
        _field(line, "PAGE", 8)
        _font(line.add_run(" / "), 8, MUTED)
        _field(line, "NUMPAGES", 8)
        _top_rule(line)

    def _letterhead(self, paragraph: Any, size: float) -> None:
        letterhead = self.data["letterhead"]
        logo = self._logo()
        if logo is not None:
            try:
                paragraph.add_run().add_picture(str(logo), height=Mm(9))
                paragraph.add_run("  ")
            except Exception as error:  # noqa: BLE001 - an unreadable image is left out
                self.warnings.append(f"Word: the letterhead logo could not be added ({error}).")
        if letterhead.get("organisation"):
            run = paragraph.add_run(letterhead["organisation"])
            _font(run, size, INK, bold=True)
        for line in letterhead.get("lines", []):
            run = paragraph.add_run()
            run.add_break()
            run.text = line
            _font(run, size, MUTED)

    def _logo(self) -> Path | None:
        name = self.data["letterhead"].get("logo")
        if not name or self.root is None:
            return None
        path = self.root / name
        return path if path.is_file() else None

    # ─────────────────────────── document ───────────────────────────

    def write(self) -> None:
        cover = self.data["cover"]
        section = self.doc.sections[0]
        if cover["show"]:
            section.different_first_page_header_footer = True
            self._cover()
        self._header_footer(section)
        if self.data.get("contents"):
            heading = self.doc.add_paragraph()
            _font(heading.add_run("Contents"), 15, INK, bold=True)
            _field(self.doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', 10)
        first = True
        for s in self.data["sections"]:
            if s["landscape"]:
                self._landscape()
            elif not first or self.data.get("contents") or cover["show"]:
                self._page_break()
            first = False
            self._heading(1, s["title"])
            for block in s["blocks"]:
                self._block(block)

    def _cover(self) -> None:
        doc = self.doc
        meta = self.data["meta"]
        cover = self.data["cover"]
        if self.data["letterhead"].get("enabled"):
            self._letterhead(doc.add_paragraph(), size=9)
            doc.add_paragraph()
        else:
            for _ in range(3):
                doc.add_paragraph()
        _font(doc.add_paragraph().add_run(cover["title"].upper()), 10, MUTED)
        _font(doc.add_paragraph().add_run(meta["project"]), 24, INK, bold=True)
        if cover.get("subtitle"):
            _font(doc.add_paragraph().add_run(cover["subtitle"]), 13, MUTED)
        _bottom_rule(doc.add_paragraph(), size=8)
        rows = [*meta.get("details", []), ["Code basis", meta["code"]], ["Date", meta["date"]]]
        table = doc.add_table(rows=0, cols=2)
        for label, value in rows:
            cells = table.add_row().cells
            _cell_text(cells[0], label, 10.5, MUTED)
            _cell_text(cells[1], value, 10.5, INK)
        _widths(table, [Mm(45), Mm(115)])
        doc.add_paragraph()
        status = self.data["status"]
        if cover.get("status") and status.get("headline"):
            state = status.get("state")
            fill = {"fail": FAIL_FILL, "pass": PASS_FILL}.get(state, WARN_FILL)
            color = {"fail": FAIL_INK, "pass": PASS_INK}.get(state, INK)
            box = doc.add_paragraph()
            _font(box.add_run(status["headline"]), 11, color, bold=True)
            _shade(box, fill)
        sign = meta.get("titleBlock") or []
        if cover.get("titleBlock") and sign:
            doc.add_paragraph()
            table = doc.add_table(rows=0, cols=4)
            table.style = "Table Grid"
            cells: list[Any] = []
            for label, value in sign:
                if not cells:
                    cells = list(table.add_row().cells)
                cell = cells.pop(0)
                _cell_text(cell, label, 7.5, MUTED)
                cell.add_paragraph()
                _font(cell.paragraphs[-1].add_run(value), 9, INK)
        if self.data["warnings"]:
            box = doc.add_paragraph()
            _font(box.add_run("Read before use. "), 9, INK, bold=True)
            _font(
                box.add_run(
                    f"This report carries {len(self.data['warnings'])} warning(s), listed in the "
                    "Limitations section."
                ),
                9,
                INK,
            )
            _shade(box, WARN_FILL)
        if cover.get("note"):
            _font(doc.add_paragraph().add_run(cover["note"]), 9.5, INK)
        tail = doc.add_paragraph()
        _font(tail.add_run(f"Prepared with SAEngine {meta['version']} (AGPL-3.0). "), 9, MUTED)
        _font(tail.add_run(meta["units"]), 9, MUTED)
        _font(
            doc.add_paragraph().add_run(
                "This report presents the results of computer analysis and design. It shall be "
                "reviewed, checked and signed by the structural engineer responsible for the "
                "design."
            ),
            9,
            MUTED,
        )

    def _page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _landscape(self) -> None:
        if not self.appendix:
            self.appendix = True
            self.numbers = [0, 0, 0]
        current = self.doc.sections[-1]
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = current.page_height, current.page_width
        section.different_first_page_header_footer = False

    def _heading(self, level: int, text: str) -> None:
        self.numbers[level - 1] += 1
        for deeper in range(level, 3):
            self.numbers[deeper] = 0
        first = chr(ord("A") + self.numbers[0] - 1) if self.appendix else str(self.numbers[0])
        number = ".".join([first, *map(str, self.numbers[1:level])])
        self.doc.add_heading(f"{number}  {text}", level=level)

    # ─────────────────────────── blocks ───────────────────────────

    def _block(self, b: dict[str, Any]) -> None:
        kind = b["t"]
        doc = self.doc
        if kind == "h":
            self._heading(min(int(b["level"]), 3), b["text"])
        elif kind == "p":
            paragraph = doc.add_paragraph()
            size = b.get("size") or 10
            color = MUTED if b.get("muted") else INK
            if b.get("strong"):
                _font(paragraph.add_run(b["strong"] + " "), size, color, bold=True)
            _font(paragraph.add_run(b["text"]), size, color)
        elif kind == "pairs":
            table = doc.add_table(rows=0, cols=2)
            for label, value in b["rows"]:
                cells = table.add_row().cells
                _cell_text(cells[0], label, 9, MUTED)
                _cell_text(cells[1], value, 9, INK)
                _rule_below(cells)
            share = float(str(b.get("width", "55%")).rstrip("%")) / 100
            _widths(table, [Mm(170 * share), Mm(170 * (1 - share))])
            doc.add_paragraph()
        elif kind == "table":
            self._table(b["head"], b["rows"], b["align"], b.get("bad"), b["size"])
        elif kind == "checks":
            self._checks(b)
        elif kind == "bullets":
            for item in b["items"]:
                paragraph = doc.add_paragraph(style="List Bullet")
                _font(paragraph.add_run(item), b.get("size", 9), INK)
        elif kind == "note":
            paragraph = doc.add_paragraph()
            if b.get("strong"):
                _font(paragraph.add_run(b["strong"] + " "), 9, INK, bold=True)
            _font(paragraph.add_run(b["text"]), 9, INK)
            _shade(paragraph, WARN_FILL)
        elif kind == "empty":
            run = doc.add_paragraph().add_run(b["text"])
            _font(run, 9, MUTED, italic=True)
        elif kind == "label":
            tone = b.get("tone")
            color = FAIL_INK if tone == "fail" else MUTED if tone == "muted" else INK
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            _font(paragraph.add_run(b["text"]), 9, color, bold=tone != "muted")
        else:
            self.warnings.append(f"Word: a block of kind {kind!r} was not drawn.")

    def _table(
        self,
        head: list[str],
        rows: list[list[str]],
        align: list[str],
        bad: list[bool] | None,
        size: float,
    ) -> Any:
        table = self.doc.add_table(rows=1, cols=len(head))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (cell, text) in enumerate(zip(table.rows[0].cells, head, strict=False)):
            _cell_text(cell, text, size, INK, bold=True, align=align[i] if i < len(align) else "l")
            _cell_fill(cell, HEAD_FILL)
        _repeat_header(table.rows[0])
        for r, row in enumerate(rows):
            failing = bool(bad and r < len(bad) and bad[r])
            cells = table.add_row().cells
            for i, text in enumerate(row[: len(head)]):
                fail_cell = failing and text == "FAIL"
                _cell_text(
                    cells[i],
                    text,
                    size,
                    FAIL_INK if fail_cell else INK,
                    bold=fail_cell,
                    align=align[i] if i < len(align) else "l",
                )
                if failing:
                    _cell_fill(cells[i], FAIL_FILL)
            _rule_below(cells)
        self.doc.add_paragraph()
        return table

    def _checks(self, b: dict[str, Any]) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        _font(paragraph.add_run(b["title"]), 9, INK if b["ok"] else FAIL_INK, bold=True)
        table = self.doc.add_table(rows=1, cols=7)
        for i, (cell, text) in enumerate(zip(table.rows[0].cells, CHECK_HEAD, strict=False)):
            _cell_text(cell, text, 7.5, INK, bold=True, align=CHECK_ALIGN[i])
            _cell_fill(cell, HEAD_FILL)
        _repeat_header(table.rows[0])
        for r, row in enumerate(b["rows"]):
            failing = bool(r < len(b["fails"]) and b["fails"][r])
            cells = table.add_row().cells
            for i in range(7):
                text = row[i] if i < len(row) else ""
                fail_cell = failing and i == 6
                _cell_text(
                    cells[i], text, 7.5, FAIL_INK if fail_cell else INK, bold=fail_cell,
                    align=CHECK_ALIGN[i],
                )  # fmt: skip
                if failing:
                    _cell_fill(cells[i], FAIL_FILL)
            remark = row[7] if len(row) > 7 else ""
            if remark:
                merged = table.add_row().cells
                cell = merged[0].merge(merged[6])
                _cell_text(cell, remark, 7.5, MUTED)
                if failing:
                    _cell_fill(cell, FAIL_FILL)
            _rule_below(table.rows[-1].cells)
        self.doc.add_paragraph()


# ─────────────────────────────────────── helpers ───────────────────────────────────────


def _font(run: Any, size: float, color: RGBColor, bold: bool = False, italic: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def _cell_text(
    cell: Any,
    text: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
    align: str = "l",
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    _font(paragraph.add_run(str(text)), size, color, bold=bold)


def _cell_fill(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _rule_below(cells: Any) -> None:
    for cell in cells:
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "2")
        bottom.set(qn("w:color"), RULE)
        borders.append(bottom)
        cell._tc.get_or_add_tcPr().append(borders)


def _repeat_header(row: Any) -> None:
    """The header row repeats on every page the table runs onto."""
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def _widths(table: Any, widths: list[Any]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=False):
            cell.width = width


def _shade(paragraph: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _border(paragraph: Any, side: str, size: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    line = OxmlElement(f"w:{side}")
    line.set(qn("w:val"), "single")
    line.set(qn("w:sz"), str(size))
    line.set(qn("w:space"), "1")
    line.set(qn("w:color"), RULE if size < 8 else "141414")
    borders.append(line)


def _bottom_rule(paragraph: Any, size: int = 4) -> None:
    _border(paragraph, "bottom", size)


def _top_rule(paragraph: Any, size: int = 4) -> None:
    _border(paragraph, "top", size)


def _field(paragraph: Any, instruction: str, size: float) -> None:
    """A Word field (page number, page count, table of contents), filled in by Word."""
    run = paragraph.add_run()
    _font(run, size, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = (
        "Update the field (F9) to fill it in." if instruction.startswith("TOC") else "1"
    )
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, placeholder, end):
        run._r.append(element)
