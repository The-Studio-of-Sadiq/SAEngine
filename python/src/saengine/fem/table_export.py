"""Result tables written to an Excel workbook (``.xlsx``, openpyxl): one sheet per table.

The FEM results' tables panel sends the tables as it shows them (columns and rows of text or
numbers); each becomes a sheet with a bold, frozen heading row and its notes under the rows.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

#: Excel's limit on a sheet name, and the characters it forbids in one.
_NAME_LENGTH = 31
_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def sheet_name(title: str, taken: set[str]) -> str:
    """A valid sheet name for ``title``, unique among ``taken`` (which it joins)."""
    base = _FORBIDDEN.sub(" ", title).strip().strip("'") or "Sheet"
    name, number = base[:_NAME_LENGTH], 2
    while name.lower() in taken:
        suffix = f" ({number})"
        name, number = base[: _NAME_LENGTH - len(suffix)] + suffix, number + 1
    taken.add(name.lower())
    return name


def write_xlsx(path: str | Path, sheets: list[dict[str, Any]]) -> dict[str, Any]:
    """Writes ``sheets`` (``[{"name", "columns", "rows", "notes"?}]``) to ``path``:
    ``{"path", "sheets"}``."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError(f"{target.name}: an Excel workbook must end in .xlsx")
    book = Workbook()
    book.remove(book.active)
    taken: set[str] = set()
    for table in sheets or [{"name": "Results", "columns": [], "rows": []}]:
        sheet = book.create_sheet(sheet_name(str(table.get("name", "")), taken))
        columns = [str(c) for c in table.get("columns", [])]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in table.get("rows", []):
            sheet.append([c if isinstance(c, (int, float)) else str(c) for c in row])
        notes = table.get("notes") or []
        if notes:
            sheet.append([])
            for note in notes:
                sheet.append([str(note)])
        sheet.freeze_panes = "A2"
        for k, title in enumerate(columns, 1):
            width = max(
                [
                    len(title),
                    *(len(str(r[k - 1])) for r in table.get("rows", [])[:200] if len(r) >= k),
                ]
            )
            sheet.column_dimensions[sheet.cell(1, k).column_letter].width = min(40, width + 2)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return {"path": str(target), "sheets": len(book.sheetnames)}
