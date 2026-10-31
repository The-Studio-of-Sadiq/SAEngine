"""The model as an Excel workbook, and back: grids, stories, members, sections and loads.

Engineers and architects edit long lists faster in a spreadsheet. :func:`export_workbook` writes
the floor plan (``05 Floor Plan/floor-plan.json``), the object types (``04 Objects/objects.json``,
the sections) and the member loads (``05 Floor Plan/loads.json``) to one ``.xlsx``, a sheet per
list; :func:`import_workbook` reads an edited workbook back, validates it with the same models and
returns the new files with a diff of what changed, for the floor plan to preview before it writes.

There is no second schema: each sheet's columns are the fields of the library model of its rows
(:mod:`saengine.library.floorplan`, :mod:`saengine.library.objects`,
:mod:`saengine.library.line_loads`). A nested model's fields get dotted columns (``at.x``,
``wall.height``); lists and choices of models (a slab's corners, a section's profile) are one
cell of JSON. An empty cell is the field's default.

What the workbook leaves out stays as it is: the underlays, and the footings, which are placed in
9 Substructure (the floor plan shows them). Units are the files' own: mm, degrees, kN/m, kN.
"""

from __future__ import annotations

import json
import types
from dataclasses import dataclass
from pathlib import Path
from typing import Annotated, Any, Literal, Union, get_args, get_origin
from zipfile import BadZipFile

from pydantic import BaseModel, TypeAdapter, ValidationError

from saengine.library import floorplan, line_loads, objects

#: Files the workbook covers, by the key the RPC uses for their text.
FLOOR_PLAN, OBJECTS, LOADS = "floorPlan", "objects", "loads"
README_SHEET = "Read me"
#: Rows past this are not read (a sheet pasted over by mistake).
MAX_ROWS = 20_000


@dataclass(frozen=True)
class Column:
    """One column: the field's path in the row's model and how its cells read."""

    path: tuple[str, ...]
    #: ``"scalar"`` (a number, text, yes/no or choice) or ``"json"`` (anything else).
    kind: str
    text: bool
    title: str
    description: str

    @property
    def key(self) -> str:
        return ".".join(self.path)


@dataclass(frozen=True)
class Sheet:
    name: str
    file: str
    #: Where the rows are in the file's JSON (``("grids", "x")`` and so on).
    location: tuple[str, ...]
    #: The type of one row (a model, or a discriminated union of models).
    row_type: Any
    #: The models whose fields are the columns (all members of a union).
    models: tuple[type[BaseModel], ...]
    #: A fixed column first, e.g. ``axis`` of the grid lines.
    fixed: tuple[str, str] | None = None
    #: The field naming a row (``id``, or ``label`` for grid lines).
    key: str = "id"


def _strip(annotation: Any) -> Any:
    """``annotation`` without ``Annotated`` and without ``None`` in a union."""
    while get_origin(annotation) is Annotated:
        annotation = get_args(annotation)[0]
    if get_origin(annotation) in (Union, types.UnionType):
        rest = [a for a in get_args(annotation) if a is not type(None)]
        if len(rest) == 1:
            return _strip(rest[0])
    return annotation


def _item_type(model: type[BaseModel], name: str) -> Any:
    """The element type of the list field ``name`` of ``model``."""
    annotation = model.model_fields[name].annotation
    return get_args(annotation)[0]


def _union_models(annotation: Any) -> tuple[type[BaseModel], ...]:
    annotation = _strip(annotation)
    if isinstance(annotation, type) and issubclass(annotation, BaseModel):
        return (annotation,)
    return tuple(m for a in get_args(annotation) for m in _union_models(a))


SHEETS: tuple[Sheet, ...] = (
    Sheet(
        "Grids",
        FLOOR_PLAN,
        ("grids",),
        floorplan.GridLine,
        (floorplan.GridLine,),
        ("axis", "X or Y"),
        "label",
    ),
    Sheet("Stories", FLOOR_PLAN, ("stories",), floorplan.Story, (floorplan.Story,)),
    Sheet("Columns", FLOOR_PLAN, ("columns",), floorplan.PlacedColumn, (floorplan.PlacedColumn,)),
    Sheet("Beams", FLOOR_PLAN, ("beams",), floorplan.PlacedBeam, (floorplan.PlacedBeam,)),
    Sheet("Walls", FLOOR_PLAN, ("walls",), floorplan.PlacedWall, (floorplan.PlacedWall,)),
    Sheet("Slabs", FLOOR_PLAN, ("slabs",), floorplan.PlacedSlab, (floorplan.PlacedSlab,)),
    Sheet(
        "Sections",
        OBJECTS,
        ("objects",),
        _item_type(objects.ObjectLibrary, "objects"),
        _union_models(_item_type(objects.ObjectLibrary, "objects")),
    ),
    Sheet("Wall types", LOADS, ("wallTypes",), line_loads.WallType, (line_loads.WallType,)),
    Sheet(
        "Member loads",
        LOADS,
        ("loads",),
        _item_type(line_loads.MemberLoads, "loads"),
        _union_models(_item_type(line_loads.MemberLoads, "loads")),
    ),
)


# ── columns from the models ──


def _scalar(annotation: Any) -> tuple[bool, bool]:
    """Whether a field is one cell of plain value, and whether that value is text."""
    annotation = _strip(annotation)
    if annotation in (str, int, float, bool):
        return True, annotation is str
    if get_origin(annotation) is Literal:
        return True, all(isinstance(v, str) for v in get_args(annotation))
    if get_origin(annotation) in (Union, types.UnionType):
        parts = [_scalar(a) for a in get_args(annotation) if a is not type(None)]
        if parts and all(p[0] for p in parts):
            return True, any(p[1] for p in parts) and not any(
                _strip(a) in (int, float) for a in get_args(annotation)
            )
    return False, False


def _model_columns(model: type[BaseModel], prefix: tuple[str, ...] = ()) -> list[Column]:
    out: list[Column] = []
    for name, info in model.model_fields.items():
        if name == "schema_":
            continue
        key = info.alias or name
        path = (*prefix, key)
        inner = _strip(info.annotation)
        if isinstance(inner, type) and issubclass(inner, BaseModel):
            out.extend(_model_columns(inner, path))
            continue
        scalar, text = _scalar(info.annotation)
        out.append(
            Column(
                path,
                "scalar" if scalar else "json",
                text,
                info.title or name,
                info.description or "",
            )
        )
    return out


def columns(sheet: Sheet) -> list[Column]:
    """The sheet's columns: the fields of its models, the discriminator (``kind``) first."""
    seen: dict[str, Column] = {}
    for model in sheet.models:
        for column in _model_columns(model):
            seen.setdefault(column.key, column)
    ordered = sorted(seen.values(), key=lambda c: (c.key != "kind", c.key != sheet.key))
    if sheet.fixed:
        name, description = sheet.fixed
        ordered.insert(0, Column((name,), "scalar", True, name, description))
    return ordered


# ── rows ──


def _get(data: Any, path: tuple[str, ...]) -> Any:
    for key in path:
        if not isinstance(data, dict):
            return None
        data = data.get(key)
    return data


def _cell(value: Any, column: Column) -> Any:
    if value is None:
        return None
    if column.kind == "json":
        return json.dumps(value, ensure_ascii=False, separators=(", ", ": "))
    return value


def flatten(item: dict[str, Any], cols: list[Column]) -> dict[str, Any]:
    """A row's cells by column key (``None`` for empty)."""
    return {c.key: _cell(_get(item, c.path), c) for c in cols}


def _number_text(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else repr(value)


def _read_cell(value: Any, column: Column) -> Any:
    """A cell's value as the field takes it; ``None`` for an empty cell."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if column.kind == "json":
        if not isinstance(value, str):
            return value
        try:
            return json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{column.key}: not valid JSON ({exc.msg})") from exc
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in ("true", "yes"):
            return True
        if value.lower() in ("false", "no"):
            return False
        return value
    if column.text and isinstance(value, int | float) and not isinstance(value, bool):
        # Excel turns a label typed as 1 into the number 1.
        return _number_text(value)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def unflatten(cells: dict[str, Any], cols: list[Column]) -> dict[str, Any]:
    """The row's item; empty cells are left out (the field's default)."""
    item: dict[str, Any] = {}
    for column in cols:
        value = _read_cell(cells.get(column.key), column)
        if value is None:
            continue
        target = item
        for key in column.path[:-1]:
            target = target.setdefault(key, {})
        target[column.path[-1]] = value
    return item


def _rows_of(sheet: Sheet, doc: dict[str, Any]) -> list[dict[str, Any]]:
    """The sheet's items in a file's JSON (grid lines with their ``axis``)."""
    if sheet.fixed and sheet.location == ("grids",):
        grids = doc.get("grids") or {}
        return [
            {"axis": axis.upper(), **line}
            for axis in ("x", "y")
            for line in grids.get(axis) or []
            if isinstance(line, dict)
        ]
    items = _get(doc, sheet.location)
    return [i for i in items or [] if isinstance(i, dict)]


def _parse_doc(text: str | None) -> dict[str, Any]:
    if not text or not text.strip():
        return {}
    value = json.loads(text)
    return value if isinstance(value, dict) else {}


# ── export ──


def export_workbook(
    out_path: str | Path,
    floor_plan_text: str,
    objects_text: str | None = None,
    loads_text: str | None = None,
) -> dict[str, Any]:
    """Writes the model to ``out_path``: ``{"path", "sheets": {name: rows}}``."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    docs = {
        FLOOR_PLAN: _parse_doc(floor_plan_text),
        OBJECTS: _parse_doc(objects_text),
        LOADS: _parse_doc(loads_text),
    }
    book = Workbook()
    readme = book.active
    readme.title = README_SHEET
    for line in _README:
        readme.append([line])
    readme.column_dimensions["A"].width = 110
    counts: dict[str, int] = {}
    for sheet in SHEETS:
        ws = book.create_sheet(sheet.name)
        cols = columns(sheet)
        ws.append([c.key for c in cols])
        for i, column in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True)
            note = column.title + (f": {column.description}" if column.description else "")
            if column.kind == "json":
                note += " (JSON)"
            cell.comment = Comment(note, "SAEngine")
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(column.key) + 4))
        rows = _rows_of(sheet, docs[sheet.file])
        for item in rows:
            cells = flatten(item, cols)
            ws.append([cells[c.key] for c in cols])
        for i, column in enumerate(cols, start=1):
            if column.text or column.kind == "json":
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    row[0].number_format = "@"
        ws.freeze_panes = "B2"
        counts[sheet.name] = len(rows)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    return {"path": str(out), "sheets": counts}


_README = (
    "SAEngine model workbook: grids, stories, members, sections and member loads.",
    "",
    (
        "Edit the sheets and import the workbook back from the floor plan (More Actions > Import "
        "from Excel...). SAEngine checks every row and shows what changes before it writes."
    ),
    (
        "Row 1 of each sheet holds the field names: do not rename them. Hover a heading for its "
        "meaning and unit."
    ),
    (
        "An empty cell takes the field's default. A column marked (JSON) holds a list or a choice "
        "of shapes as JSON text."
    ),
    (
        "Adding a row adds an item; deleting a row deletes it; a row is matched by its id (grid "
        "lines by their label)."
    ),
    (
        "Units: lengths, positions and elevations in mm; angles in degrees; line loads in kN/m; "
        "point loads in kN; unit weights in kN/m3."
    ),
    (
        "Grids: axis X (vertical lines, labelled A, B, C...) or Y (horizontal lines, labelled 1, "
        "2, 3...)."
    ),
    (
        "Members are placed on the grid: a column 'at' a grid point (at.x, at.y, at.dx, at.dy); "
        "a beam or wall from start to end; a slab by its corners."
    ),
    (
        "Footings and underlays are not in the workbook: footings are placed in 9 Substructure, "
        "and both stay as they are on import."
    ),
)


# ── import ──


@dataclass
class _Read:
    items: list[dict[str, Any]]
    problems: list[str]


def _read_sheet(ws: Any, sheet: Sheet) -> _Read:
    cols = columns(sheet)
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    problems: list[str] = []
    if header is None:
        return _Read([], problems)
    names = [str(h).strip() if h is not None else "" for h in header]
    known = {c.key for c in cols}
    for name in names:
        if name and name not in known:
            problems.append(
                f"{sheet.name}: unknown column '{name}' (row 1 must keep the field names)"
            )
    adapter = TypeAdapter(sheet.row_type)
    items: list[dict[str, Any]] = []
    for number, values in enumerate(rows, start=2):
        if number > MAX_ROWS:
            problems.append(f"{sheet.name}: rows past {MAX_ROWS} are not read")
            break
        if values is None or all(
            v is None or (isinstance(v, str) and not v.strip()) for v in values
        ):
            continue
        cells = {name: value for name, value in zip(names, values, strict=False) if name in known}
        try:
            item = unflatten(cells, cols)
        except ValueError as exc:
            problems.append(f"{sheet.name} row {number}: {exc}")
            continue
        check = {k: v for k, v in item.items() if k != "axis"} if sheet.fixed else item
        if sheet.fixed and str(item.get("axis", "")).upper() not in ("X", "Y"):
            problems.append(f"{sheet.name} row {number}: axis must be X or Y")
            continue
        try:
            adapter.validate_python(check)
        except ValidationError as exc:
            for error in exc.errors():
                where = ".".join(str(p) for p in error["loc"] if not isinstance(p, int))
                problems.append(f"{sheet.name} row {number}: {where or 'row'}: {error['msg']}")
            continue
        items.append(item)
    ids = [str(i.get(sheet.key)) for i in items]
    for duplicate in sorted({i for i in ids if ids.count(i) > 1}):
        problems.append(f"{sheet.name}: '{duplicate}' is on more than one row")
    return _Read(items, problems)


def _normal(value: Any) -> Any:
    """``value`` for comparison: no nulls or empty objects, whole floats as ints."""
    if isinstance(value, dict):
        out = {k: _normal(v) for k, v in value.items()}
        return {k: v for k, v in out.items() if v is not None and v != {}}
    if isinstance(value, list):
        return [_normal(v) for v in value]
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _canonical(item: dict[str, Any], row_type: Any) -> Any:
    """The item as its model reads it, for comparison (defaults filled in)."""
    try:
        model = TypeAdapter(row_type).validate_python(item)
        return _normal(TypeAdapter(row_type).dump_python(model, mode="json", by_alias=True))
    except ValidationError:
        return _normal(item)


def _clean(value: Any) -> Any:
    if isinstance(value, dict):
        return {k: _clean(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_clean(v) for v in value]
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _text(doc: dict[str, Any]) -> str:
    """The file's text as the workbench writes it (tab indent, newline at the end)."""
    return json.dumps(doc, indent="\t", ensure_ascii=False) + "\n"


def _merge(
    sheet: Sheet, old: list[dict[str, Any]], new: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """The new list (an unchanged item keeps its old JSON exactly) and the sheet's diff."""
    cols = columns(sheet)
    key = sheet.key

    def ident(item: dict[str, Any]) -> str:
        prefix = f"{item.get('axis', '')}:" if sheet.fixed else ""
        return prefix + str(item.get(key))

    def strip_fixed(item: dict[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in item.items() if not (sheet.fixed and k == sheet.fixed[0])}

    before = {ident(i): i for i in old}
    after = {ident(i): i for i in new}
    merged: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    for name, item in after.items():
        previous = before.get(name)
        if previous is not None and _canonical(strip_fixed(previous), sheet.row_type) == _canonical(
            strip_fixed(item), sheet.row_type
        ):
            merged.append(previous)
            continue
        merged.append(_clean(item))
        if previous is not None:
            a, b = flatten(previous, cols), flatten(_clean(item), cols)
            fields = [
                {"field": c.key, "old": a[c.key], "new": b[c.key]}
                for c in cols
                if _normal(_read_cell(a[c.key], c)) != _normal(_read_cell(b[c.key], c))
            ]
            changed.append({"id": str(item.get(key)), "fields": fields})
    diff = {
        "added": [str(after[n].get(key)) for n in after if n not in before],
        "removed": [str(before[n].get(key)) for n in before if n not in after],
        "changed": changed,
    }
    return merged, diff


def import_workbook(
    path: str | Path,
    floor_plan_text: str,
    objects_text: str | None = None,
    loads_text: str | None = None,
) -> dict[str, Any]:
    """Reads an edited workbook against the current files.

    Returns ``{"ok", "problems", "diff": {sheet: {added, removed, changed}}, "texts": {file: text},
    "changed": [file, ...]}``; ``ok`` is false (and ``texts`` empty) when a row or a whole file
    does not validate. Sheets missing from the workbook leave their lists as they are.
    """
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        return {
            "ok": False,
            "problems": [f"{Path(path).name} is not a readable .xlsx workbook: {exc}"],
            "diff": {},
            "texts": {},
            "changed": [],
        }
    docs = {
        FLOOR_PLAN: _parse_doc(floor_plan_text),
        OBJECTS: _parse_doc(objects_text),
        LOADS: _parse_doc(loads_text),
    }
    new_docs = {name: json.loads(json.dumps(doc)) for name, doc in docs.items()}
    problems: list[str] = []
    diff: dict[str, Any] = {}
    try:
        for sheet in SHEETS:
            if sheet.name not in book.sheetnames:
                continue
            read = _read_sheet(book[sheet.name], sheet)
            problems.extend(read.problems)
            old = _rows_of(sheet, docs[sheet.file])
            merged, sheet_diff = _merge(sheet, old, read.items)
            if not any(sheet_diff.values()):
                continue  # the list stays exactly as it is in the file
            diff[sheet.name] = sheet_diff
            target = new_docs[sheet.file]
            if sheet.fixed and sheet.location == ("grids",):
                target["grids"] = {
                    axis.lower(): [
                        {k: v for k, v in line.items() if k != "axis"}
                        for line in merged
                        if str(line.get("axis", "")).upper() == axis
                    ]
                    for axis in ("X", "Y")
                }
            else:
                target[sheet.location[0]] = merged
    finally:
        book.close()

    changed = [
        name
        for name in (FLOOR_PLAN, OBJECTS, LOADS)
        if _normal(new_docs[name]) != _normal(docs[name])
    ]
    if not problems:
        problems.extend(_check(new_docs, changed))
    if problems:
        return {"ok": False, "problems": problems, "diff": diff, "texts": {}, "changed": changed}
    for name in changed:
        new_docs[name].setdefault("schema", _SCHEMAS[name])
    return {
        "ok": True,
        "problems": [],
        "diff": diff,
        "texts": {name: _text(new_docs[name]) for name in changed},
        "changed": changed,
    }


_SCHEMAS = {
    FLOOR_PLAN: floorplan.SCHEMA_ID,
    OBJECTS: objects.SCHEMA_ID,
    LOADS: line_loads.SCHEMA_ID,
}


def _check(docs: dict[str, dict[str, Any]], changed: list[str]) -> list[str]:
    """The whole files' checks: references between rows, sheets and files."""
    problems: list[str] = []

    def validate(model: type[BaseModel], name: str, label: str) -> BaseModel | None:
        if not docs[name] and name not in changed:
            return None
        try:
            return model.model_validate(docs[name])
        except ValidationError as exc:
            for error in exc.errors():
                where = ".".join(str(p) for p in error["loc"])
                message = error["msg"].removeprefix("Value error, ")
                problems.append(f"{label}: {where + ': ' if where else ''}{message}")
            return None

    plan = validate(floorplan.FloorPlan, FLOOR_PLAN, "Floor plan")
    library = validate(objects.ObjectLibrary, OBJECTS, "Sections")
    loads = validate(line_loads.MemberLoads, LOADS, "Member loads")
    if isinstance(plan, floorplan.FloorPlan) and isinstance(library, objects.ObjectLibrary):
        problems.extend(f"Floor plan: {p}" for p in plan.check_objects(library))
    if isinstance(plan, floorplan.FloorPlan) and isinstance(loads, line_loads.MemberLoads):
        stories = {s.id for s in plan.stories}
        members = {m.id for m in plan.beams} | {m.id for m in plan.columns}
        for load in loads.loads:
            if load.story not in stories:
                problems.append(f"Member loads: {load.id}: no story '{load.story}'")
            if load.member not in members:
                problems.append(f"Member loads: {load.id}: no beam or column '{load.member}'")
    return problems


def sheet_names() -> list[str]:
    return [s.name for s in SHEETS]


__all__ = [
    "SHEETS",
    "columns",
    "export_workbook",
    "flatten",
    "import_workbook",
    "sheet_names",
    "unflatten",
]
