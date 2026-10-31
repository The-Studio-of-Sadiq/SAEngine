"""The bill of quantities as an Excel workbook (``.xlsx``, openpyxl), the costing and the
embodied carbon with it.

Sheets:

- ``Summary``: the cost per part, per group and in total, and the embodied carbon, as formulas
  over the sheets below (so a rate changed in Excel carries through);
- ``Bill of Quantities``: every item, its breakdown by element grouped under it (Excel's
  outline: collapse to the items);
- ``Costing``: every item with the rate item that prices it, ``amount = quantity × conversion ×
  rate`` as a formula, its share per part;
- ``Embodied Carbon``: kg of each material, its factor and source, ``kg CO2e = kg × factor``;
- ``Schedule of Rates``: the schedule used, by chapter;
- ``Notes``: the bill's and the costing's notes.

Cement stays in kg and in 50 kg bags and sand and stone chips in cft, as the bill measures
them (:mod:`saengine.boq.quantities`).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from saengine.boq.costing import PARTS, CarbonFactors, RateSchedule, cost_bill

QUANTITY = "#,##0.000"
MONEY = "#,##0.00"
FACTOR = "0.000"

GROUP_TITLES = {
    "earthwork": "Earthwork",
    "concrete": "Concrete",
    "masonry": "Masonry",
    "finishes": "Finishes",
    "reinforcement": "Reinforcement",
    "formwork": "Formwork",
    "steel": "Structural steel",
    "mep": "MEP",
    "materials": "Materials",
}


def _widths(sheet: Any, widths: list[int]) -> None:
    from openpyxl.utils import get_column_letter

    for k, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(k)].width = width


def _heading(sheet: Any, columns: list[str]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet.append(columns)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
    sheet.freeze_panes = sheet.cell(sheet.max_row + 1, 1)


def write_boq_xlsx(
    path: str | Path,
    bill: dict[str, Any],
    rates: RateSchedule | None = None,
    factors: CarbonFactors | None = None,
    title: str = "",
    sample_rates: bool = False,
) -> dict[str, Any]:
    """Writes the bill (``boq.json`` as a dict), priced by ``rates`` with the carbon of
    ``factors`` when given, to ``path``: ``{"path", "sheets", "total", "carbonTotal"}``."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError(f"{target.name}: an Excel workbook must end in .xlsx")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    costing = cost_bill(bill, rates, factors) if rates is not None else None

    book = Workbook()
    summary = book.active
    summary.title = "Summary"

    # ── Bill of Quantities ──
    sheet = book.create_sheet("Bill of Quantities")
    sheet.sheet_properties.outlinePr.summaryBelow = False
    _heading(sheet, ["Code", "Group", "Description", "Element", "Story", "Detail", "Quantity",
                     "Unit"])  # fmt: skip
    for item in bill.get("items", []):
        sheet.append([item.get("code"), GROUP_TITLES.get(item.get("group"), item.get("group")),
                      item.get("description"), None, None, None, item.get("quantity"),
                      item.get("unit")])  # fmt: skip
        row = sheet.max_row
        for cell in sheet[row]:
            cell.font = bold
        sheet.cell(row, 7).number_format = QUANTITY
        for entry in item.get("breakdown") or []:
            sheet.append([None, None, None, entry.get("element"), entry.get("story"),
                          entry.get("detail"), entry.get("quantity"), item.get("unit")])  # fmt: skip
            sheet.cell(sheet.max_row, 7).number_format = QUANTITY
            sheet.row_dimensions[sheet.max_row].outlineLevel = 1
    _widths(sheet, [9, 15, 48, 14, 10, 34, 14, 7])

    # ── Costing ──
    cost_parts: list[str] = []
    cost_last = 1
    if costing is not None and rates is not None:
        sheet = book.create_sheet("Costing")
        parts = [p for p in PARTS if any(p in li.byPart for li in costing.lines)]
        columns = ["Code", "Group", "Description", "Quantity", "Unit", "Rate item",
                   "Description of rate item", "Rate unit", "Units per bill unit",
                   f"Rate ({rates.currency})", f"Amount ({rates.currency})", *parts, "Status"]  # fmt: skip
        _heading(sheet, columns)
        for line in costing.lines:
            sheet.append([line.code, line.group, line.description, line.quantity, line.unit,
                          line.rateCode, line.rateDescription, line.rateUnit, line.conversion,
                          line.rate])  # fmt: skip
            r = sheet.max_row
            if line.conversion is not None:
                sheet.cell(r, 11, f"=IF(ISNUMBER(J{r}),D{r}*I{r}*J{r},0)")
            else:
                sheet.cell(r, 11, 0)
            share_total = sum(line.byPart.values()) or 0.0
            for k, part in enumerate(parts):
                share = line.byPart.get(part, 0.0) / share_total if share_total else 0.0
                sheet.cell(r, 12 + k, f"=K{r}*{share:.6g}" if share else 0)
                sheet.cell(r, 12 + k).number_format = MONEY
            sheet.cell(r, 12 + len(parts), line.status)
            sheet.cell(r, 4).number_format = QUANTITY
            sheet.cell(r, 9).number_format = "0.0000"
            sheet.cell(r, 10).number_format = MONEY
            sheet.cell(r, 11).number_format = MONEY
        last = sheet.max_row
        sheet.append(["Total", None, None, None, None, None, None, None, None, None,
                      f"=SUM(K2:K{last})",
                      *[f"=SUM({get_column_letter(12 + k)}2:{get_column_letter(12 + k)}{last})"
                        for k in range(len(parts))]])  # fmt: skip
        for cell in sheet[sheet.max_row]:
            cell.font = bold
            cell.number_format = MONEY
        _widths(sheet, [9, 13, 44, 13, 6, 10, 44, 9, 10, 12, 16, *[15] * len(parts), 14])
        cost_parts = parts
        cost_last = last

    # ── Embodied Carbon ──
    carbon_last = None
    if costing is not None and factors is not None:
        sheet = book.create_sheet("Embodied Carbon")
        _heading(sheet, ["Material", "Mass (kg)", "Factor (kg CO2e/kg)", "Carbon (kg CO2e)",
                         "Carbon (t CO2e)", "Bill items", "Source"])  # fmt: skip
        for line in costing.carbon:
            sheet.append([line.material, line.mass, line.factor])
            r = sheet.max_row
            sheet.cell(r, 4, f"=IF(ISNUMBER(C{r}),B{r}*C{r},0)")
            sheet.cell(r, 5, f"=D{r}/1000")
            sheet.cell(r, 6, ", ".join(line.items))
            sheet.cell(r, 7, line.source).alignment = wrap
            sheet.cell(r, 2).number_format = QUANTITY
            sheet.cell(r, 3).number_format = FACTOR
            sheet.cell(r, 4).number_format = MONEY
            sheet.cell(r, 5).number_format = QUANTITY
        carbon_last = sheet.max_row
        sheet.append(["Total", None, None, f"=SUM(D2:D{carbon_last})", f"=SUM(E2:E{carbon_last})"])
        for cell in sheet[sheet.max_row]:
            cell.font = bold
        sheet.cell(sheet.max_row, 4).number_format = MONEY
        sheet.cell(sheet.max_row, 5).number_format = QUANTITY
        _widths(sheet, [30, 16, 18, 18, 14, 24, 90])

    # ── Schedule of Rates ──
    if rates is not None:
        sheet = book.create_sheet("Schedule of Rates")
        sheet.append([rates.name])
        sheet.cell(1, 1).font = Font(bold=True, size=13)
        if rates.edition:
            sheet.append([rates.edition])
        for note in rates.notes:
            sheet.append([note])
        sheet.append([])
        _heading(sheet, ["Item No.", "Description of item", "Unit", f"Rate ({rates.currency})"])
        chapter = None
        for item in rates.items:
            ch = rates.chapter_of(item.code)
            if ch is not None and ch is not chapter:
                chapter = ch
                sheet.append([ch.code, ch.title.upper()])
                for cell in sheet[sheet.max_row]:
                    cell.font = bold
            sheet.append([item.code, item.description, item.unit, item.rate])
            sheet.cell(sheet.max_row, 2).alignment = wrap
            sheet.cell(sheet.max_row, 4).number_format = MONEY
        _widths(sheet, [10, 80, 8, 14])

    # ── Notes ──
    notes = [*bill.get("notes", []), *bill.get("problems", [])]
    if costing is not None:
        notes += costing.notes
    if notes:
        sheet = book.create_sheet("Notes")
        _widths(sheet, [140])
        for note in notes:
            sheet.append([str(note)])
            sheet.cell(sheet.max_row, 1).alignment = wrap

    # ── Summary ──
    summary.append([title or "Bill of Quantities"])
    summary.cell(1, 1).font = Font(bold=True, size=14)
    if costing is not None and rates is not None:
        summary.append([f"Rates: {rates.name}" + (f" ({rates.edition})" if rates.edition else "")])
        if sample_rates:
            summary.append([("The rates are the SAMPLE's made-up example rates, not real ones: "
                            "enter your own in 10 BOQ/rates.json.")])  # fmt: skip
            summary.cell(summary.max_row, 1).font = Font(bold=True, color="C00000")
        summary.append([])
        summary.append(["Cost by part", f"Amount ({rates.currency})"])
        for cell in summary[summary.max_row]:
            cell.font = bold
        for k, part in enumerate(cost_parts):
            column = get_column_letter(12 + k)
            summary.append([part, f"='Costing'!{column}{cost_last + 1}"])
            summary.cell(summary.max_row, 2).number_format = MONEY
        summary.append([])
        summary.append(["Cost by group", f"Amount ({rates.currency})"])
        for cell in summary[summary.max_row]:
            cell.font = bold
        for group in dict.fromkeys(li.group for li in costing.lines):
            formula = f"=SUMIF('Costing'!B2:B{cost_last},\"{group}\",'Costing'!K2:K{cost_last})"
            summary.append([GROUP_TITLES.get(group, group), formula])
            summary.cell(summary.max_row, 2).number_format = MONEY
        summary.append(["Total", f"='Costing'!K{cost_last + 1}"])
        for cell in summary[summary.max_row]:
            cell.font = bold
        summary.cell(summary.max_row, 2).number_format = MONEY
        priced = sum(1 for li in costing.lines if li.status == "priced")
        summary.append([f"Priced items: {priced} of {len(costing.lines)} (see Costing, Status)"])
    if carbon_last is not None and costing is not None:
        summary.append([])
        summary.append(["Embodied carbon (materials, production only)", "t CO2e"])
        for cell in summary[summary.max_row]:
            cell.font = bold
        for part, carbon in costing.carbonParts.items():
            summary.append([part, carbon / 1e3])
            summary.cell(summary.max_row, 2).number_format = QUANTITY
        summary.append(["Total", f"='Embodied Carbon'!E{carbon_last + 1}"])
        for cell in summary[summary.max_row]:
            cell.font = bold
        summary.cell(summary.max_row, 2).number_format = QUANTITY
    _widths(summary, [46, 20])

    book.calculation.fullCalcOnLoad = True
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return {
        "path": str(target),
        "sheets": len(book.sheetnames),
        "total": costing.total if costing else None,
        "carbonTotal": costing.carbonTotal if costing else None,
    }


def write_boq_xlsx_texts(
    path: str | Path,
    boq_text: str | None,
    rates_text: str | None = None,
    carbon_text: str | None = None,
    title: str = "",
) -> dict[str, Any]:
    """``write_boq_xlsx`` of the project files' texts (the shipped sample rates and factors when
    not given): ``{"path", ...}``, or ``{"problems"}``."""
    from saengine.boq.costing import CARBON_FILE, RATES_FILE, parse_carbon, parse_rates
    from saengine.pipeline.run_all import parse_boq

    bill = parse_boq(boq_text)
    if bill is None:
        return {"problems": ["Measure the bill of quantities (10 BOQ) first."]}
    problems = []
    rates = factors = None
    try:
        rates = parse_rates(rates_text)
    except ValueError as error:
        problems.append(f"{RATES_FILE}: {error}")
    try:
        factors = parse_carbon(carbon_text)
    except ValueError as error:
        problems.append(f"{CARBON_FILE}: {error}")
    if problems:
        return {"problems": problems}
    sample = not (rates_text and rates_text.strip())
    return write_boq_xlsx(path, bill, rates, factors, title, sample_rates=sample)
