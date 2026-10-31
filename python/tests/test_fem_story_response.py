"""Storey response, member detail and table export of the FEM results (saengine.fem.story_response,
saengine.fem.member_detail, saengine.fem.table_export)."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pytest
from test_fem_derive import OBJECTS, PLAN, SETTINGS

from saengine.fem import member_detail
from saengine.fem.derive import derive_fem
from saengine.fem.story_response import (
    _overturning,
    _shears,
    _storey_loads,
    storey_levels,
    story_response,
)
from saengine.fem.table_export import sheet_name, write_xlsx
from saengine.library import default_material_library
from saengine.library.bim import generate

# ─────────────────────────────── pure helpers ───────────────────────────────


def test_storey_forces_shears_and_overturning_by_hand() -> None:
    levels = [0.0, 3.0, 6.0, 9.0]
    # Loads near the levels go to the nearest one; a load at the base carries no storey.
    loads = [(3.0, 10.0, 0.0), (5.99, 20.0, -5.0), (9.0, 30.0, 0.0), (0.0, 7.0, 0.0)]
    fx, fy = _storey_loads(loads, levels)
    assert fx.tolist() == [7.0, 10.0, 20.0, 30.0]
    assert fy.tolist() == [0.0, 0.0, -5.0, 0.0]
    # Storey shears: the forces at and above each storey's top level.
    assert _shears(fx).tolist() == [60.0, 50.0, 30.0]
    # Overturning at each level from the forces above it.
    assert _overturning(fx, levels).tolist() == pytest.approx(
        [10 * 3 + 20 * 6 + 30 * 9, 20 * 3 + 30 * 6, 30 * 3, 0.0]
    )


def test_chain_orders_a_member_from_one_end() -> None:
    class E:
        def __init__(self, i: int, j: int) -> None:
            self.i, self.j = i, j

    a, b, c = E(2, 3), E(4, 3), E(1, 2)  # 1-2, 2-3, 3-4 given out of order, one reversed
    order = member_detail._chain([a, b, c])
    # Walked from one free end to the other (member_forces then turns it to run upwards).
    walk = [(e.j, e.i) if back else (e.i, e.j) for e, back in order]
    assert all(walk[k][1] == walk[k + 1][0] for k in range(len(walk) - 1))
    assert {walk[0][0], walk[-1][1]} == {1, 4}
    assert member_detail.member_of("C3#2") == "C3" and member_detail.member_of("B1") == "B1"


def test_demand_points_keep_what_matters_per_combination() -> None:
    points = [
        (100e3, 10e6, 1e6, "1.4D", "C1 at 0.00 L"),
        (90e3, 30e6, 2e6, "1.4D", "C1 at 1.00 L"),
        (95e3, 5e6, 40e6, "1.4D", "C1 at 0.50 L"),
        (50e3, 1e6, 1e6, "0.9D + EX", "C1 at 0.00 L"),
    ]
    out = member_detail._demand_points(points)
    by = {(p["combination"], p["where"]) for p in out}
    assert by == {
        ("1.4D", "C1 at 0.00 L"),
        ("1.4D", "C1 at 1.00 L"),
        ("1.4D", "C1 at 0.50 L"),
        ("0.9D + EX", "C1 at 0.00 L"),
    }
    first = next(p for p in out if p["where"] == "C1 at 1.00 L")
    assert (first["P"], first["Mx"], first["My"]) == pytest.approx((90.0, 30.0, 2.0))


def test_column_interaction_is_worked_out_once_per_run(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from saengine.design import rc, required

    out = tmp_path / "08 FEM" / "model.results"
    out.mkdir(parents=True)
    (out / "manifest.json").write_text("{}", encoding="utf-8")
    model = tmp_path / "08 FEM" / "model.osmodel"
    model.write_text("{}", encoding="utf-8")

    class Section:
        phi_pn_max, phi_tension = 2000e3, 500e3

    class Surface:
        def __init__(self, section: object) -> None:
            pass

        def moment_capacity(self, pu: float, angle: float) -> float:
            return 100e6 * (1 - abs(pu - 750e3) / 1250e3)

        def check(self, pu: float, mux: float, muy: float) -> rc.ColumnCheck:
            return rc.ColumnCheck(mux / 100e6, 100e6, mux, "")

        def contour(self, pu: float) -> np.ndarray:
            return np.array([[100e6, 0], [0, 100e6], [-100e6, 0], [0, -100e6]])

    calls = []

    def gather(out_dir: object, project_path: object, substructure: bool = True) -> object:
        calls.append(1)
        member = {
            "id": "C1",
            "kind": "column",
            "type": "C400",
            "section": Section(),
            "points": [(300e3, 60e6, 5e6, "1.4D", "C1 at 0.00 L")],
        }

        class Gathered:
            def __init__(self) -> None:
                self.collect = {"members": [member, {"id": "B1", "kind": "beam", "points": None}]}
                self.assumed = {"C400"}
                # A column type with assumed bars: the design report cannot be told from these.
                self.plain = None

        return Gathered()

    monkeypatch.setattr(required, "gather", gather)
    monkeypatch.setattr(rc, "InteractionSurface", Surface)
    first = member_detail.column_interaction(out, model, "C1")
    column = first["columns"]["C1"]
    assert column["assumedBars"] and column["phiPnMax"] == pytest.approx(2000.0)
    assert column["governing"]["utilisation"] == pytest.approx(0.6)
    assert len(column["curves"]["x"]) == member_detail.CURVE_POINTS
    assert column["contour"][0] == pytest.approx([100.0, 0.0])
    # Read again from the file beside the results, until the run changes.
    assert member_detail.column_interaction(out, model)["columns"].keys() == {"C1"}
    assert len(calls) == 1
    (out / "manifest.json").write_text('{"again": 1}', encoding="utf-8")
    assert member_detail.column_interaction(out, model, "B1")["columns"] == {}
    assert len(calls) == 2


def test_tables_to_excel(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    taken: set[str] = set()
    assert (
        sheet_name("Storey drifts: 1.2D + 1.6L [X/Y]", taken) == "Storey drifts  1.2D + 1.6L  X Y"
    )
    assert sheet_name("Storey drifts: 1.2D + 1.6L [X/Y]", taken).endswith(" (2)")
    path = tmp_path / "out" / "results.xlsx"
    written = write_xlsx(
        path,
        [
            {"name": "Joint displacements", "columns": ["Node", "ux [mm]"], "rows": [[1, 2.5]]},
            {
                "name": "Reactions",
                "columns": ["Node", "Fz [kN]"],
                "rows": [[1, -10]],
                "notes": ["n"],
            },
        ],
    )
    assert written["sheets"] == 2
    book = load_workbook(path)
    assert book.sheetnames == ["Joint displacements", "Reactions"]
    assert [c.value for c in book["Joint displacements"][2]] == [1, 2.5]
    assert book["Reactions"]["A4"].value == "n"
    with pytest.raises(ValueError):
        write_xlsx(tmp_path / "results.csv", [])


# ─────────────────────────────── a solved bay ───────────────────────────────


@pytest.fixture(scope="module")
def solved(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path, dict]:
    root = tmp_path_factory.mktemp("bay")
    model_dir = root / "08 FEM"
    model_dir.mkdir()
    building = generate(PLAN, "s")
    (root / "06 BIM").mkdir()
    (root / "06 BIM" / "bim.json").write_text(building.model_dump_json(), encoding="utf-8")
    (root / "02 Project Settings").mkdir()
    (root / "02 Project Settings" / "project-settings.json").write_text(
        SETTINGS.model_dump_json(), encoding="utf-8"
    )
    d = derive_fem(building, OBJECTS, default_material_library(), SETTINGS, None, model_dir)
    model = model_dir / "model.osmodel"
    model.write_text(json.dumps(d.project), encoding="utf-8")
    out = model_dir / "model.results"
    cases = {a["name"]: a["id"] for a in d.project["analyses"]}
    wanted = [n for n in cases if n.split(":")[0] in ("D", "EX", "Modal", "RX")]
    run = subprocess.run(
        [sys.executable, "-m", "opensees_studio.run", "--project", str(model)]
        + ["--cases", *(str(cases[n]) for n in wanted), "--out", str(out)],
        check=False,
        capture_output=True,
        text=True,
        timeout=600,
    )
    assert run.returncode == 0, run.stderr[-2000:]
    return out, model, d.project


def test_storey_response_of_a_one_storey_bay(solved: tuple[Path, Path, dict]) -> None:
    out, model, project = solved
    response = story_response(out, model)
    levels = response["levels"]
    assert len(levels) == 2 and levels[0]["z"] == pytest.approx(0.0)
    assert levels == storey_levels(project, model)
    height = levels[1]["z"] - levels[0]["z"]
    by_label = {r["label"].split(":")[0]: r for r in response["results"]}

    # EX: the storey shear is the sum of the pattern's loads, and the supports take it back.
    pattern = next(p for p in project["load_patterns"] if p["name"] == "EX")
    total = sum(load["forces"][0] for load in pattern["nodal_loads"])
    ex = by_label["EX"]
    top = ex["storeys"][1]
    assert top["shearX"] == pytest.approx(total, rel=1e-9)
    assert top["forceX"] == pytest.approx(total, rel=1e-9)
    assert ex["baseReaction"][0] == pytest.approx(-total, rel=1e-3)
    assert ex["storeys"][0]["overturningX"] == pytest.approx(total * height, rel=1e-9)
    # One storey on fixed supports: the drift ratio is the floor's displacement over its height.
    assert top["driftX"] == pytest.approx(abs(top["dispX"]) / height, rel=1e-6)
    assert top["shearY"] == pytest.approx(0.0, abs=1e-6)

    # Dead load: no lateral force, the supports carry the weight.
    dead = by_label["D"]
    assert dead["storeys"][1]["shearX"] == pytest.approx(0.0, abs=1e-9)
    assert dead["baseReaction"][2] > 0

    # The bay sways as one mass: mode 1 carries (nearly) all the mass in X or Y.
    modal = response["modal"]
    assert modal["required"] == 0.9 and modal["modes"][0]["period"] > 0
    assert max(modal["modes"][-1]["sumUX"], modal["modes"][-1]["sumUY"]) == pytest.approx(
        1.0, abs=0.05
    )
    # The response spectrum's storey shear: of one mass, about M_eff·Sa of the governing mode.
    rx = by_label["RX"]
    assert rx["kind"] == "spectrum" and rx["storeys"][1]["shearX"] > 0
    assert response["spectrum"]["periods"] and response["driftLimit"]["ratio"] == 0.02


def test_spectrum_charts_share_the_design_drift_and_scale(solved: tuple[Path, Path, dict]) -> None:
    from saengine.fem import seismic_design as rs

    out, model, project = solved
    response = story_response(out, model)
    rx = next(r for r in response["results"] if r["label"].startswith("RX"))
    # Without the spectrum design of this run the forces are unscaled, and said to be.
    assert rx["forceScale"] is None and response["spectrumDesign"] is None
    assert any("unscaled" in n for n in response["notes"])
    # The drift is the spectrum design's: the modal drifts combined, not the peaks' difference.
    levels = [lv["z"] for lv in response["levels"]]
    design = rs.storey_drifts(project, rs._spectra(out), levels, 1.0, float("inf"))
    x = next(d for d in design if d["direction"] == "X")
    assert rx["storeys"][1]["driftX"] == pytest.approx(x["ratio"], rel=1e-9)

    # With it (written for this run), the storey forces and shears are at the design level.
    summary = out / rs.RS_FOLDER / rs.SUMMARY_FILE
    summary.parent.mkdir(exist_ok=True)
    drift = [{**x, "ratio": 0.01, "allowableRatio": 0.02, "ok": True}]
    summary.write_text(
        json.dumps(
            {
                "schema": rs.SUMMARY_SCHEMA,
                "available": True,
                "scale": {"x": 2.0, "y": 3.0, "clause": rs.CLAUSES["scale"]},
                "drift": drift,
                "run": {"stamp": rs._stamp(out, model, False)},
            }
        ),
        encoding="utf-8",
    )
    try:
        scaled = story_response(out, model)
    finally:
        summary.unlink()
    sx = next(r for r in scaled["results"] if r["label"].startswith("RX"))
    assert sx["forceScale"] == 2.0
    assert sx["storeys"][1]["shearX"] == pytest.approx(2 * rx["storeys"][1]["shearX"])
    assert sx["storeys"][0]["overturningX"] == pytest.approx(2 * rx["storeys"][0]["overturningX"])
    assert sx["storeys"][1]["driftX"] == rx["storeys"][1]["driftX"]  # drifts stay unscaled
    assert scaled["spectrumDesign"]["scale"] == {"X": 2.0, "Y": 3.0}
    assert scaled["spectrumDesign"]["drift"][0]["ratio"] == 0.01


def test_member_forces_of_a_column(solved: tuple[Path, Path, dict]) -> None:
    out, model, project = solved
    column = next(e["name"] for e in project["elements"] if e.get("name", "").startswith("C"))
    detail = member_detail.member_forces(out, model, member_detail.member_of(column))
    assert detail["kind"] == "column" and detail["stations"][0] == 0.0
    section_id = next(e["section_id"] for e in project["elements"] if e["name"] == column)
    assert detail["section"] == next(
        s["name"] for s in project["sections"] if s["id"] == section_id
    )
    assert detail["stations"][-1] == pytest.approx(detail["length"])
    dead = next(r for r in detail["results"] if r["label"].startswith("D"))
    n = [f[0] for f in dead["forces"]]
    assert max(n) < 0  # the dead load compresses the column
    # The deflection is measured from the chord: zero at both ends.
    assert dead["deflection"][0] == pytest.approx([0, 0, 0], abs=1e-12)
    assert dead["deflection"][-1] == pytest.approx([0, 0, 0], abs=1e-12)
    with pytest.raises(ValueError):
        member_detail.member_forces(out, model, "nope")


# ───────────────────── member axes, result tables and their export ─────────────────────


def test_member_axes_put_a_beams_gravity_bending_in_m3() -> None:
    from saengine.fem.derive import local_axes, member_axes, to_member_axes

    # A beam along X: the analysis gives it local z up (gravity bends it about local y, My).
    beam = np.array(local_axes(np.zeros(3), np.array([6.0, 0, 0])))
    assert beam[2] == pytest.approx([0, 0, 1])
    shown = member_axes(beam)
    assert shown[1] == pytest.approx([0, 0, 1])  # axis 2 up
    assert np.cross(shown[0], shown[1]) == pytest.approx(shown[2])  # right-handed
    # Sagging (tension below: My < 0) is M3 > 0; the vertical shear Vz is V2.
    forces = np.array([10.0, 0.0, -25.0, 0.0, -20.0, 0.0])  # N Vy Vz T My Mz
    assert to_member_axes(forces, beam) == pytest.approx([10, -25, 0, 0, 0, 20])
    # A column keeps its axes (the design and the 3D view already read it so).
    column = np.array(local_axes(np.zeros(3), np.array([0, 0, 3.0])))
    assert member_axes(column) == pytest.approx(column)


def test_beam_forces_are_shown_in_member_axes(solved: tuple[Path, Path, dict]) -> None:
    from saengine.fem.display import result_view

    out, model, project = solved
    dead = next(a["id"] for a in project["analyses"] if a["name"].startswith("D:"))
    view = result_view(out, model, f"case:{dead}")
    beams = [e for e in view["elements"] if e["name"].startswith("B") and "forces" in e]
    assert beams
    for element in beams:
        f = np.asarray(element["forces"])
        assert element["y"] == pytest.approx([0, 0, 1], abs=1e-6)  # axis 2 up
        # Gravity: V2 and M3 carry it, V3 and M2 (sideways) nothing (N-X1: it was the other way).
        assert np.abs(f[:, 1]).max() > 1e3 and np.abs(f[:, 5]).max() > 1e3
        assert np.abs(f[:, 2]).max() < 1.0 and np.abs(f[:, 4]).max() < 1.0
    beam = beams[0]["name"].split("#")[0]
    detail = member_detail.member_forces(out, model, beam)
    d = next(r for r in detail["results"] if r["label"].startswith("D"))
    m3 = [row[5] for row in d["forces"]]
    assert max(abs(v) for v in m3) > 1e3 and max(abs(row[4]) for row in d["forces"]) < 1.0


def test_result_tables_one_at_a_time(solved: tuple[Path, Path, dict]) -> None:
    from saengine.fem import result_tables as rt

    out, model, _ = solved
    index = rt.table_index(out, model)["results"]
    dead = next(r for r in index if r["label"].startswith("D:"))
    assert [t["id"] for t in dead["tables"]] == ["displacements", "reactions", "members"]
    modal = next(r for r in index if r["label"].startswith("Modal"))
    assert [t["id"] for t in modal["tables"]] == ["modal:0"]
    # The bay ran D, EX and Modal: a combination needing L is listed, not formed.
    combination = next(r for r in index if r["group"] == "combination")
    assert combination["tables"] == [] and "not formed" in combination["warnings"][0]

    members = rt.result_table(out, model, dead["key"], "members")
    assert members["columns"][:3] == ["Member", "Station [m]", "P [kN]"]
    names = {row[0] for row in members["rows"]}
    assert all("#" not in n for n in names)  # members, not analysis pieces
    beam_rows = [r for r in members["rows"] if r[0].startswith("B")]
    assert max(abs(r[7]) for r in beam_rows) > 1 and max(abs(r[6]) for r in beam_rows) < 0.01
    # kN with fixed decimals; a beam's stations run from 0 to its length.
    assert all(round(v, rt.FORCE_DECIMALS) == v for r in members["rows"] for v in r[2:])
    first = [r for r in beam_rows if r[0] == beam_rows[0][0]]
    assert first[0][1] == 0.0 and first[-1][1] > first[0][1]
    # The analysis's own forces, in kN: the station's V2 is the element's vertical shear.
    reactions = rt.result_table(out, model, dead["key"], "reactions")
    total = reactions["rows"][-1]
    assert total[0] == "Total" and total[3] > 0  # the supports carry the weight
    weight = sum(r[3] for r in reactions["rows"][:-1])
    assert total[3] == pytest.approx(weight, abs=0.05)
    frequencies = rt.result_table(out, model, modal["key"], "modal:0")
    assert frequencies["rows"]


def test_tables_export_streams_to_a_workbook(
    solved: tuple[Path, Path, dict], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import load_workbook

    from saengine.fem import result_tables as rt

    out, model, _ = solved
    path = tmp_path / "results.xlsx"
    extra = [{"name": "Storeys", "columns": ["Level"], "rows": [["Base"]]}]
    written = rt.export_xlsx(out, model, path, tables=["reactions", "members"], extra=extra)
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == ["Support reactions", "Member forces", "Storeys"]
    assert written["sheets"] == 3 and written["rows"] > 0
    heading = next(book["Member forces"].iter_rows(max_row=1, values_only=True))
    assert heading[:3] == ("Result", "Member", "Station [m]")
    book.close()
    assert not list(tmp_path.glob("*.partial.xlsx"))

    # Stopped: no workbook at the path, and no half-written one beside it.
    stopped = tmp_path / "stopped.xlsx"
    real = rt.result_table

    def stop_after_one(*args: object) -> dict:
        rt.cancel_export()
        return real(*args)

    monkeypatch.setattr(rt, "result_table", stop_after_one)
    with pytest.raises(rt.ExportCancelled):
        rt.export_xlsx(out, model, stopped, tables=["reactions", "members"])
    assert not stopped.exists() and not list(tmp_path.glob("*.partial.xlsx"))
    assert rt.export_progress()["running"] is False


def test_combination_labels_show_every_case() -> None:
    from saengine.fem.combinations import combination_label
    from saengine.library.settings import LoadCombination

    present = {"D", "SD", "L", "EX"}
    # The D factor applies to SD too: a name of factors is written out from them.
    assert (
        combination_label(
            LoadCombination(name="1.2D + 1.0EX + 1.0L", D=1.2, EX=1.0, L=1.0), present
        )
        == "1.2D + 1.2SD + 1.0L + 1.0EX"
    )
    assert combination_label(LoadCombination(name="0.9D - 1.0EX", D=0.9, EX=-1.0), present) == (
        "0.9D + 0.9SD − 1.0EX"
    )
    # A name of the engineer's own is kept.
    assert combination_label(LoadCombination(name="Service", D=1.0, L=1.0), present) == "Service"
