"""The BOQ's costing by a schedule of rates and its embodied carbon (hand-checked), and the Excel
workbook of the bill."""

from __future__ import annotations

from pathlib import Path

import pytest

from saengine.boq.costing import (
    CFT,
    SQFT,
    CarbonFactors,
    RateSchedule,
    conversion,
    cost_bill,
    cost_texts,
    parse_carbon,
    parse_rates,
)
from saengine.boq.quantities import SAND_DENSITY


def _item(code, group, quantity, unit, *, kind=None, diameter=None, breakdown=None, **extra):
    return {"code": code, "group": group, "description": f"{group} {code}", "quantity": quantity,
            "unit": unit, "kind": kind, "diameter": diameter,
            "breakdown": breakdown or [], **extra}  # fmt: skip


BILL = {
    "schema": "saengine.boq/1",
    "items": [
        _item(
            "E.01",
            "earthwork",
            10.0,
            "m3",
            kind="excavation",
            breakdown=[{"element": "F1", "quantity": 10.0}],
        ),
        _item(
            "C.01",
            "concrete",
            2.0,
            "m3",
            kind="footing",
            breakdown=[{"element": "F1", "quantity": 2.0}],
        ),
        _item(
            "C.02",
            "concrete",
            3.0,
            "m3",
            kind="column",
            breakdown=[{"element": "C1", "quantity": 3.0}],
        ),
        _item(
            "R.01",
            "reinforcement",
            100.0,
            "kg",
            diameter=16,
            breakdown=[{"element": "F1", "quantity": 40.0}, {"element": "C1", "quantity": 60.0}],
        ),
        _item(
            "MT.01",
            "materials",
            500.0,
            "kg",
            kind="cement",
            breakdown=[{"element": "C1", "quantity": 500.0}],
        ),
        _item(
            "MT.02",
            "materials",
            10.0,
            "bag",
            kind="bags",
            breakdown=[{"element": "C1", "quantity": 10.0}],
        ),
        _item(
            "MT.03",
            "materials",
            35.3147,
            "cft",
            kind="sand",
            breakdown=[{"element": "C1", "quantity": 35.3147}],
        ),
        _item("P.01", "mep", 4.0, "nr", kind="fixture"),
    ],
}

RATES = RateSchedule.model_validate({
    "name": "test", "currency": "BDT",
    "chapters": [{"code": "03", "title": "Concrete"}],
    "items": [
        {"code": "02.01", "description": "dig", "unit": "cft", "rate": 10,
         "match": {"group": "earthwork"}},
        {"code": "03.01", "description": "footing", "unit": "m3", "rate": 1000,
         "match": {"group": "concrete", "kind": ["footing"]}},
        {"code": "03.02", "description": "other", "unit": "m3", "rate": 2000,
         "match": {"group": "concrete"}},
        {"code": "05.01", "description": "bars", "unit": "t", "rate": 100_000,
         "match": {"group": "reinforcement", "diameter": [12.01, 40]}},
        {"code": "10.01", "description": "bags", "unit": "bag", "rate": None,
         "match": {"group": "materials", "kind": ["bags"]}},
        {"code": "10.02", "description": "wrong", "unit": "m2", "rate": 1,
         "match": {"group": "materials", "kind": ["cement"]}},
    ],
})  # fmt: skip

FACTORS = CarbonFactors.model_validate({
    "factors": [
        {"id": "cement", "material": "Cement", "factor": 0.5,
         "match": {"group": "materials", "kind": ["cement"]}},
        {"id": "rebar", "material": "Rebar", "factor": 2.0, "match": {"group": "reinforcement"}},
        {"id": "sand", "material": "Sand", "factor": 0.01,
         "match": {"group": "materials", "kind": ["sand"]}},
        {"id": "brick", "material": "Brick", "factor": None, "match": {"group": "concrete"}},
    ],
})  # fmt: skip


def test_unit_conversions() -> None:
    assert conversion("m3", "m3") == 1.0
    assert conversion("cft", "m3") == pytest.approx(CFT)  # 1 m³ = 35.3147 cft
    assert conversion("m3", "cft") == pytest.approx(1 / CFT)
    assert conversion("sft", "m2") == pytest.approx(SQFT)  # 1 m² = 10.7639 sft
    assert conversion("t", "kg") == pytest.approx(1e-3)
    assert conversion("m2", "kg") is None


def test_costing_by_item_part_and_total() -> None:
    costing = cost_bill(BILL, RATES, FACTORS)
    lines = {li.code: li for li in costing.lines}
    # 10 m³ dug at 10 per cft: 10 × 35.3147 × 10.
    assert lines["E.01"].amount == pytest.approx(10 * CFT * 10)
    assert lines["E.01"].byPart == {"Substructure": pytest.approx(10 * CFT * 10)}
    # The first matching rate item wins: the footing at 1000, the column at 2000.
    assert lines["C.01"].rateCode == "03.01" and lines["C.01"].amount == pytest.approx(2000)
    assert lines["C.02"].rateCode == "03.02" and lines["C.02"].amount == pytest.approx(6000)
    # 100 kg of 16 mm bars at 100 000 per t; F1 is a footing, C1 a column.
    assert lines["R.01"].amount == pytest.approx(10_000)
    assert lines["R.01"].byPart == {
        "Substructure": pytest.approx(4000), "Superstructure": pytest.approx(6000)
    }  # fmt: skip
    assert lines["MT.02"].status == "not rated" and lines["MT.02"].amount is None
    assert lines["MT.01"].status == "unit mismatch"
    assert lines["P.01"].status == "no rate item"
    assert costing.total == pytest.approx(10 * CFT * 10 + 2000 + 6000 + 10_000)
    assert costing.parts["Substructure"] == pytest.approx(10 * CFT * 10 + 2000 + 4000)
    assert costing.parts["Superstructure"] == pytest.approx(6000 + 6000)
    assert costing.groups["concrete"] == pytest.approx(8000)
    assert any("No rate item" in n for n in costing.notes)


def test_embodied_carbon() -> None:
    costing = cost_bill(BILL, RATES, FACTORS)
    carbon = {li.id: li for li in costing.carbon}
    assert carbon["cement"].carbon == pytest.approx(500 * 0.5)  # bags not counted again
    assert carbon["rebar"].carbon == pytest.approx(100 * 2.0)
    # 35.3147 cft = 1 m³ of loose sand at the take-off's density.
    assert carbon["sand"].mass == pytest.approx(SAND_DENSITY)
    assert carbon["sand"].carbon == pytest.approx(SAND_DENSITY * 0.01)
    assert carbon["brick"].carbon is None  # no factor: listed, not counted
    assert costing.carbonTotal == pytest.approx(250 + 200 + SAND_DENSITY * 0.01)
    assert costing.carbonParts["Substructure"] == pytest.approx(40 * 2.0)
    assert any("no factor for Brick" in n for n in costing.notes)


def test_shipped_data_parse() -> None:
    rates = parse_rates(None)
    assert rates.items and all(r.code for r in rates.items)
    assert "made up" in " ".join(rates.notes).lower()
    factors = parse_carbon(None)
    assert {f.id for f in factors.factors} >= {"cement", "rebar", "steel"}
    assert all(f.source for f in factors.factors)


def test_cost_texts_needs_a_bill() -> None:
    assert "problems" in cost_texts(None)
    result = cost_texts('{"schema": "saengine.boq/1", "items": []}')
    assert result["sampleRates"] is True and result["costing"]["total"] == 0


def test_workbook(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from saengine.boq.workbook import write_boq_xlsx

    out = write_boq_xlsx(tmp_path / "boq.xlsx", BILL, RATES, FACTORS, "Test")
    assert out["total"] == pytest.approx(10 * CFT * 10 + 18_000)
    book = load_workbook(out["path"])
    assert book.sheetnames == ["Summary", "Bill of Quantities", "Costing", "Embodied Carbon",
                               "Schedule of Rates", "Notes"]  # fmt: skip
    costing = book["Costing"]
    assert costing["K2"].value == "=IF(ISNUMBER(J2),D2*I2*J2,0)"
    bill = book["Bill of Quantities"]
    assert bill.row_dimensions[3].outlineLevel == 1  # E.01's breakdown under it
