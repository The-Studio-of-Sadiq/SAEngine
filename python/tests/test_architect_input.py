"""Architect input: a PDF page as a floor-plan underlay, and the model workbook round trip."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject

from saengine.library import pdf_underlay, workbook
from saengine.library.floorplan import FloorPlan, Underlay

SAMPLE = Path(__file__).resolve().parents[2] / "workbench/resources/saengine/samples/Sample Frame"


def write_pdf(
    path: Path, content: bytes, size: tuple[float, float] = (842, 595), rotate: int = 0
) -> Path:
    writer = PdfWriter()
    page = writer.add_blank_page(*size)
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = writer._add_object(stream)
    if rotate:
        page.rotate(rotate)
    writer.write(path)
    return path


# ── PDF underlay ──


def test_a_vector_page_reads_in_paper_mm(tmp_path: Path) -> None:
    # A 100 pt line under a 2x scale from (10, 10): 200 pt long = 70.556 mm.
    pdf = write_pdf(tmp_path / "plan.pdf", b"q 2 0 0 2 10 10 cm 0 0 m 100 0 l S Q")
    page = pdf_underlay.read_page(pdf)
    assert page.width == pytest.approx(842 * 25.4 / 72)
    assert page.height == pytest.approx(595 * 25.4 / 72)
    [(points, closed)] = page.polylines
    assert not closed
    assert points[0] == pytest.approx((10 * 25.4 / 72, 10 * 25.4 / 72))
    assert points[1][0] - points[0][0] == pytest.approx(200 * 25.4 / 72)


def test_rectangles_close_and_curves_are_flattened(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "plan.pdf", b"20 20 30 30 re f 0 0 m 10 10 20 0 30 10 c S")
    page = pdf_underlay.read_page(pdf)
    rect, curve = page.polylines
    assert rect[1] and len(rect[0]) == 4
    assert len(curve[0]) == 1 + pdf_underlay.CURVE_STEPS
    assert curve[0][-1] == pytest.approx((30 * 25.4 / 72, 10 * 25.4 / 72))


def test_a_clipping_path_is_not_drawn(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "plan.pdf", b"0 0 100 100 re W n 0 0 m 50 0 l S")
    assert len(pdf_underlay.read_page(pdf).polylines) == 1


def test_a_rotated_page_is_read_as_shown(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "plan.pdf", b"0 0 m 100 0 l S", size=(200, 100), rotate=90)
    page = pdf_underlay.read_page(pdf)
    assert (page.width, page.height) == pytest.approx((100 * 25.4 / 72, 200 * 25.4 / 72))
    [(points, _)] = page.polylines
    # Turned clockwise: the line along the page's bottom runs down its left edge.
    assert points[0] == pytest.approx((0, 200 * 25.4 / 72))
    assert points[1] == pytest.approx((0, 100 * 25.4 / 72))


def test_convert_caches_the_dxf_and_snap_segments(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "plan.pdf", b"0 0 m 100 0 l 100 50 l S")
    first = pdf_underlay.convert(pdf)
    assert first.vector and first.pages == 1 and not first.cached
    assert len(first.segments) == 2
    assert first.dxf.parent.name == pdf_underlay.CACHE_FOLDER
    assert "PDF-LINES" in first.dxf.read_text(encoding="utf-8")
    again = pdf_underlay.convert(pdf)
    assert again.cached and again.segments == first.segments


def test_a_page_without_paths_is_reported(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "scan.pdf", b"")
    result = pdf_underlay.convert(pdf)
    assert not result.vector
    assert "no vector linework" in result.warnings[0]


def test_a_missing_page_is_refused(tmp_path: Path) -> None:
    pdf = write_pdf(tmp_path / "plan.pdf", b"0 0 m 1 1 l S")
    with pytest.raises(pdf_underlay.PdfError, match="page 3"):
        pdf_underlay.read_page(pdf, 3)


def test_the_underlay_takes_a_pdf_page_and_a_lock() -> None:
    underlay = Underlay(story="L1", file="Underlays/plan.pdf", page=2, scale=100, locked=True)
    assert underlay.page == 2 and underlay.locked
    with pytest.raises(ValueError):
        Underlay(story="L1", file="Underlays/plan.png")
    with pytest.raises(ValueError):
        Underlay(story="L1", file="plan.pdf", page=0)


# ── model workbook ──


def sample_texts() -> tuple[str, str]:
    return (
        (SAMPLE / "05 Floor Plan/floor-plan.json").read_text(encoding="utf-8"),
        (SAMPLE / "04 Objects/objects.json").read_text(encoding="utf-8"),
    )


def test_columns_come_from_the_models() -> None:
    sheet = next(s for s in workbook.SHEETS if s.name == "Columns")
    keys = [c.key for c in workbook.columns(sheet)]
    assert keys[:4] == ["id", "type", "at.x", "at.y"]
    slabs = next(s for s in workbook.SHEETS if s.name == "Slabs")
    boundary = next(c for c in workbook.columns(slabs) if c.key == "boundary")
    assert boundary.kind == "json"
    sections = next(s for s in workbook.SHEETS if s.name == "Sections")
    assert workbook.columns(sections)[0].key == "kind"


def test_an_unedited_workbook_changes_nothing(tmp_path: Path) -> None:
    plan, objects = sample_texts()
    out = tmp_path / "model.xlsx"
    exported = workbook.export_workbook(out, plan, objects)
    assert exported["sheets"]["Grids"] == 5
    result = workbook.import_workbook(out, plan, objects)
    assert result["ok"], result["problems"]
    assert result["diff"] == {} and result["changed"] == []


def test_edits_come_back_with_a_diff(tmp_path: Path) -> None:
    plan, objects = sample_texts()
    out = tmp_path / "model.xlsx"
    workbook.export_workbook(out, plan, objects)
    book = load_workbook(out)
    grids = book["Grids"]
    header = [c.value for c in grids[1]]
    grids.cell(row=2, column=header.index("position") + 1).value = 150
    book["Beams"].delete_rows(2)
    book.save(out)
    result = workbook.import_workbook(out, plan, objects)
    assert result["ok"], result["problems"]
    assert result["changed"] == ["floorPlan"]
    grid_change = result["diff"]["Grids"]["changed"][0]
    assert grid_change["fields"] == [{"field": "position", "old": 0, "new": 150}]
    assert result["diff"]["Beams"]["removed"] == ["B1"]
    new_plan = FloorPlan.model_validate(json.loads(result["texts"]["floorPlan"]))
    assert new_plan.grids.x[0].position == 150
    assert "B1" not in {b.id for b in new_plan.beams}
    # What the workbook leaves out stays.
    assert len(new_plan.footings) == len(json.loads(plan).get("footings", []))


def test_bad_rows_are_reported_and_nothing_is_written(tmp_path: Path) -> None:
    plan, objects = sample_texts()
    out = tmp_path / "model.xlsx"
    workbook.export_workbook(out, plan, objects)
    book = load_workbook(out)
    beams = book["Beams"]
    header = [c.value for c in beams[1]]
    beams.cell(row=2, column=header.index("story") + 1).value = "NOPE"
    columns = book["Columns"]
    columns.cell(row=2, column=[c.value for c in columns[1]].index("rotation") + 1).value = "abc"
    book.save(out)
    result = workbook.import_workbook(out, plan, objects)
    assert not result["ok"] and result["texts"] == {}
    assert any(p.startswith("Columns row 2: rotation") for p in result["problems"])


def test_a_story_that_does_not_exist_fails_the_file_check(tmp_path: Path) -> None:
    plan, objects = sample_texts()
    out = tmp_path / "model.xlsx"
    workbook.export_workbook(out, plan, objects)
    book = load_workbook(out)
    beams = book["Beams"]
    header = [c.value for c in beams[1]]
    beams.cell(row=2, column=header.index("story") + 1).value = "NOPE"
    book.save(out)
    result = workbook.import_workbook(out, plan, objects)
    assert not result["ok"]
    assert any("no story 'NOPE'" in p for p in result["problems"])


def test_a_numeric_label_stays_text() -> None:
    sheet = next(s for s in workbook.SHEETS if s.name == "Grids")
    cols = workbook.columns(sheet)
    item = workbook.unflatten({"axis": "Y", "label": 1, "position": 6000.0}, cols)
    assert item == {"axis": "Y", "label": "1", "position": 6000}
